#!/usr/bin/env python3
"""Transcribe the supported English My BMW All journeys list over ADB.

Python 3.9+ and openpyxl. Uses the UI only, does not call BMW APIs, and does not
open grouped rows. Raw XML and original displayed values remain authoritative.
"""
import argparse
import csv
import hashlib
import datetime as dt
import json
import os
import re
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import time
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

MONTHS = {'JAN': 1, 'FEB': 2, 'MAR': 3, 'APR': 4, 'MAY': 5,
          'JUN': 6, 'JUL': 7, 'AUG': 8, 'SEP': 9, 'SEPT': 9,
          'OCT': 10, 'NOV': 11, 'DEC': 12}
FIELDS = ['month', 'list_sequence', 'week_context', 'destination',
          'displayed_date_time', 'distance_km', 'duration_displayed',
          'consumption_displayed', 'grouped_trip_count', 'start_location',
          'departure_time', 'arrival_time', 'raw_text', 'source_xml', 'capture_date']
PROVENANCE_FIELDS = ['source_capture', 'source_session_id', 'merge_conflicts']
CAPTURE_DATA_DIR = 'capture-data'
EXPORT_FORMATS = ('xlsx', 'csv', 'json', 'xml')


def configure_console_streams():
    """Keep diagnostics printable without changing a terminal's encoding."""
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, 'reconfigure', None)
        if reconfigure:
            try:
                reconfigure(errors='backslashreplace')
            except (OSError, ValueError):
                # Some wrappers expose reconfigure but are already closed or
                # otherwise immutable. Leave those wrappers to their caller.
                pass


def clean(s):
    return str(s or '').replace('\xa0', ' ').replace('\u202f', ' ').strip()


def bounds(node):
    values = re.findall(r'\d+', node.get('bounds', ''))
    return tuple(map(int, values)) if len(values) == 4 else (0, 0, 0, 0)


def parse_screen(xml, source=''):
    root = ET.fromstring(xml)
    nodes = list(root.iter('node'))
    descriptions = [clean(n.get('content-desc', '')) for n in nodes]
    if 'All journeys' not in descriptions:
        raise RuntimeError('Expected All journeys screen. Open that list in My BMW.')
    if not any(n.get('package') == 'de.bmw.connected.mobile20.row' for n in nodes):
        raise RuntimeError('Unexpected app package; refusing to operate another app.')
    month_node = next((n for n in nodes if re.fullmatch(
        r'[A-Z]{3,4} 20\d\d', clean(n.get('content-desc', '')))), None)
    if month_node is None:
        raise RuntimeError('Month header is missing; cannot safely navigate.')
    name, year = clean(month_node.get('content-desc')).split()
    month = f'{year}-{MONTHS[name]:02d}'
    rows, week, expected_count, expected_km = [], '', None, None
    for n in nodes:
        raw = n.get('content-desc', '')
        s = clean(raw)
        if re.fullmatch(r'(THIS WEEK|LAST WEEK|WEEK OF .+)', s):
            week = s
        m = re.fullmatch(r'Total trips\s+(\d+)', s)
        if m:
            expected_count = int(m[1])
        m = re.fullmatch(r'Total distance\s+([\d., ]+)\s*km', s)
        if m:
            expected_km = m[1]
        lines = s.splitlines()
        if len(lines) < 3:
            continue
        metric = re.fullmatch(r'([\d.,]+)\s*km\s*·\s*(.*?)\s*·\s*(.*)', lines[1])
        if not metric:
            continue
        group = re.search(r'(?m)^(\d+) trips?$', s)
        rows.append(dict(month=month, week_context=week,
                         destination=lines[0], displayed_date_time=lines[2],
                         distance_km=metric[1],
                         duration_displayed=metric[2], consumption_displayed=metric[3],
                         grouped_trip_count=int(group[1]) if group else 1,
                         raw_text=raw, source_xml=source,
                         start_location='', departure_time='', arrival_time=''))
    return dict(month=month, nodes=nodes, rows=rows,
                expected_count=expected_count, expected_km=expected_km,
                month_bounds=bounds(month_node))


def month_range(start, end):
    """Inclusive YYYY-MM range."""
    if not re.fullmatch(r'20\d\d-(0[1-9]|1[0-2])', start or '') or not re.fullmatch(r'20\d\d-(0[1-9]|1[0-2])', end or ''):
        raise ValueError('Months must use YYYY-MM (years 2000–2099).')
    a, b = map(lambda x: tuple(map(int, x.split('-'))), (start, end))
    if a > b:
        raise ValueError('Start month must be on or before end month.')
    result = []
    y, m = a
    while (y, m) <= b:
        result.append(f'{y:04d}-{m:02d}')
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return result


def months_for_year(year, today=None):
    today = today or dt.date.today()
    if year > today.year:
        raise ValueError(f'Year {year} is in the future; latest available month is {today:%Y-%m}.')
    if not 2000 <= year <= 2099:
        raise ValueError('Year must be between 2000 and 2099.')
    last = today.month if year == today.year else 12
    return month_range(f'{year}-01', f'{year}-{last:02d}')


def validate_requested_months(months, today=None):
    today = today or dt.date.today()
    future = [m for m in months if m > today.strftime('%Y-%m')]
    if future:
        raise ValueError(f"Future month request not allowed: {', '.join(future)}. Latest available month is {today:%Y-%m}.")
    return True


def capture_time(instant=None, timezone_name=None):
    """Return an aware capture timestamp in the configured IANA zone or local zone."""
    if timezone_name:
        try:
            zone = ZoneInfo(timezone_name)
        except (ZoneInfoNotFoundError, ValueError) as exc:
            raise ValueError(f'Unknown IANA timezone {timezone_name!r}.') from exc
    else:
        zone = dt.datetime.now().astimezone().tzinfo
    instant = instant or dt.datetime.now(dt.timezone.utc)
    if instant.tzinfo is None:
        raise ValueError('Capture timestamp must be timezone-aware.')
    return instant.astimezone(zone)


def timezone_label(value):
    return getattr(value.tzinfo, 'key', None) or value.tzname() or 'unknown'


def adb_server_address():
    configured = os.environ.get('ADB_SERVER_SOCKET', '')
    if configured.startswith('tcp:'):
        host_port = configured[4:]
        host, port = host_port.rsplit(':', 1)
        return host, int(port)
    return '127.0.0.1', int(os.environ.get('ADB_SERVER_PORT', '5037'))


def adb_server_env():
    host, port = adb_server_address()
    env = os.environ.copy()
    env['ADB_SERVER_SOCKET'] = f'tcp:{host}:{port}'
    return env


def adb_devices(serial=None):
    if not shutil.which('adb'):
        raise RuntimeError('adb is not installed or is not on PATH.')
    host, port = adb_server_address()
    try:
        with socket.create_connection((host, port), timeout=1): pass
    except OSError as e:
        raise RuntimeError(f'ADB server is not reachable at {host}:{port}. Start it yourself with `adb start-server` after reviewing the impact, then rerun. This tool never starts or restarts ADB.') from e
    try:
        result = subprocess.run(['adb', 'devices'], capture_output=True, text=True, timeout=15, check=True, env=adb_server_env())
    except (subprocess.SubprocessError, OSError) as e:
        raise RuntimeError(f'Could not query the connected ADB server: {e}')
    rows = [line.split() for line in result.stdout.splitlines()[1:] if line.strip() and not line.startswith('*')]
    if serial:
        rows = [row for row in rows if row[0] == serial]
    return rows


def screen_for_device(serial):
    remote = f'/sdcard/bmw-doctor-{os.getpid()}.xml'
    adb = ['adb', '-s', serial]
    env = adb_server_env()
    try:
        result = subprocess.run(adb + ['shell', 'uiautomator', 'dump', remote], capture_output=True, timeout=30, env=env)
        if result.returncode:
            raise RuntimeError(result.stderr.decode(errors='replace').strip() or 'UI XML dump failed')
        result = subprocess.run(adb + ['exec-out', 'cat', remote], capture_output=True, timeout=30, env=env)
        if result.returncode: raise RuntimeError('Cannot read UI XML from device.')
        data = result.stdout
        start = data.find(b'<?xml')
        if start < 0: start = data.find(b'<hierarchy')
        if start < 0: raise RuntimeError('Device returned no UI XML.')
        return parse_screen(data[start:])
    finally:
        best_effort_remote_xml_cleanup(adb, remote, env)


def best_effort_remote_xml_cleanup(adb, remote, env):
    """Remove only the exact temporary UI dump; cleanup must not mask capture errors."""
    try:
        result = subprocess.run(adb + ['shell', 'rm', '-f', remote], capture_output=True, timeout=15, env=env)
        if result.returncode:
            detail = result.stderr.decode(errors='replace').strip()
            print(f'Warning: could not remove temporary device XML {remote}{": " + detail if detail else ""}. It may remain on the phone.', file=sys.stderr)
    except (subprocess.SubprocessError, OSError) as exc:
        print(f'Warning: could not remove temporary device XML {remote}: {exc}. It may remain on the phone.', file=sys.stderr)


def normalize_row(row, row_number=None):
    """Normalize saved CSV strings and reject malformed capture rows centrally."""
    if not isinstance(row, dict):
        row_label = f' at row {row_number}' if row_number is not None else ''
        raise ValueError(f'Malformed saved entry{row_label}: expected an object.')
    result = {field: row.get(field, '') for field in FIELDS}
    label = f' at row {row_number}' if row_number is not None else ''
    month = result.get('month', '')
    if not isinstance(month, str) or not re.fullmatch(r'20\d\d-(0[1-9]|1[0-2])', month):
        raise ValueError(f'Malformed saved month{label}: {month!r}.')
    for field, minimum in (('list_sequence', 1), ('grouped_trip_count', 1)):
        value = result.get(field)
        try:
            result[field] = int(value) if value not in ('', None) else 1
        except (TypeError, ValueError):
            raise ValueError(f'Malformed saved {field}{label}: {value!r}.')
        if result[field] < minimum:
            raise ValueError(f'Saved {field}{label} must be at least {minimum}.')
    value = result.get('distance_km', '')
    if value not in ('', None):
        try:
            distance = float(str(value).replace(',', '.'))
        except ValueError:
            raise ValueError(f'Malformed saved distance_km{label}: {value!r}.')
        if distance < 0:
            raise ValueError(f'Saved distance_km{label} cannot be negative.')
        result['distance_km'] = str(value).strip()
    capture_date = result.get('capture_date', '')
    if capture_date:
        try:
            result['capture_date'] = dt.date.fromisoformat(capture_date).isoformat()
        except (ValueError, TypeError):
            raise ValueError(f'Malformed saved capture_date{label}: {capture_date!r}.')
    for key, value in row.items():
        if key not in result:
            result[key] = value
    return result


def atomic_bytes(path, data):
    path = Path(path)
    fd, temp_name = tempfile.mkstemp(prefix=path.name + '.', suffix='.tmp', dir=path.parent)
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(data); stream.flush(); os.fsync(stream.fileno())
        os.replace(temp_name, path)
    finally:
        if os.path.exists(temp_name): os.unlink(temp_name)


def atomic_json(path, value):
    atomic_bytes(path, (json.dumps(value, ensure_ascii=False, indent=2) + '\n').encode('utf-8'))


def atomic_csv(path, rows, fields=None):
    fields = fields or FIELDS
    path = Path(path)
    fd, temp_name = tempfile.mkstemp(prefix=path.name + '.', suffix='.tmp', dir=path.parent)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields, extrasaction='ignore')
            writer.writeheader(); writer.writerows(rows)
            stream.flush(); os.fsync(stream.fileno())
        os.replace(temp_name, path)
    finally:
        if os.path.exists(temp_name): os.unlink(temp_name)


def write_manifest(folder, manifest):
    manifest = dict(manifest)
    manifest['version'] = 1
    atomic_json(Path(folder) / 'manifest.json', manifest)


def read_capture_rows(folder):
    folder = Path(folder)
    path = folder / ('capture.csv' if (folder / 'capture.csv').exists() else 'journeys.csv')
    if not path.exists():
        raise ValueError(f'Missing {path}; expected capture.csv or journeys.csv.')
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or 'month' not in reader.fieldnames:
            raise ValueError(f'Malformed saved capture CSV: missing header or month column in {path}.')
        rows = []
        for number, row in enumerate(reader, 2):
            if not any(row.values()): continue
            rows.append(normalize_row(row, number))
    return rows


def capture_data_path(folder):
    """Return the new recovery subdirectory, or a legacy capture root."""
    folder = Path(folder).expanduser()
    nested = folder / CAPTURE_DATA_DIR
    if nested.is_dir() and any((nested / name).exists() for name in ('capture-state.json', 'manifest.json', 'capture.csv', 'journeys.csv')):
        return nested
    return folder


def parse_formats(value):
    if value == 'all':
        return list(EXPORT_FORMATS)
    formats = [part.strip().lower() for part in value.split(',') if part.strip()]
    if not formats or any(fmt not in EXPORT_FORMATS for fmt in formats):
        raise argparse.ArgumentTypeError('--formats must be a comma-separated selection of xlsx,csv,json,xml, or all.')
    if len(set(formats)) != len(formats):
        raise argparse.ArgumentTypeError('--formats cannot contain duplicates.')
    return formats


def cleanup_capture_data(folder, keep_capture=False, keep_debug=False):
    """Remove only recognized files in this tool's nested capture-data folder."""
    data = capture_data_path(folder)
    root = Path(folder).expanduser()
    owner = data / '.owner.json'
    if data == root or not data.is_dir() or not owner.is_file():
        return
    try: ownership = json.loads(owner.read_text(encoding='utf-8'))
    except (OSError, ValueError): return
    if not isinstance(ownership, dict) or ownership.get('kind') != 'bmw-journey-exporter-capture-data':
        return
    raw = data / 'raw'
    if keep_capture:
        if not keep_debug and raw.is_dir():
            for name in ownership.get('owned_raw_files', []):
                path = raw / Path(name).name
                if path.is_file(): path.unlink()
            try: raw.rmdir()
            except OSError: pass
        return
    for name in ('.owner.json', 'capture-state.json', 'capture.csv', 'checks.json', 'manifest.json'):
        path = data / name
        if path.is_file(): path.unlink()
    if raw.is_dir():
        for name in ownership.get('owned_raw_files', []):
            path = raw / Path(name).name
            if path.is_file(): path.unlink()
        try: raw.rmdir()
        except OSError: pass
    try: data.rmdir()
    except OSError: pass


def shell_quote(value):
    return "'" + str(value).replace("'", "'\\''") + "'"


def manual_timeout_restore_command(serial, timeout):
    return ' '.join(['adb', '-s', shell_quote(serial), 'shell', 'settings', 'put', 'system', 'screen_off_timeout', shell_quote(timeout)])


def resume_commands(output):
    quoted = shell_quote(output)
    return [f"make resume OUTPUT={quoted}", f"python3 bmw_export_journeys.py --resume --output {quoted}"]


def host_path_for_recovery(path, kind='output'):
    """Return the invoking host path supplied by the Docker launcher, if any."""
    if os.environ.get('BMW_HOST_PATHS_FROM_LAUNCHER') != '1':
        return path
    variable = 'BMW_HOST_REPORT_PATH' if kind == 'report' else 'BMW_HOST_OUTPUT_PATH'
    return os.environ.get(variable) or path


def row_identity(row):
    return clean(row.get('raw_text', ''))


def merge_capture_sets(captures, conflict='error'):
    """Merge ordered month sequences with provenance and explicit conflicts."""
    if conflict not in ('error', 'keep-first', 'keep-last'):
        raise ValueError('Conflict policy must be error, keep-first, or keep-last.')
    by_month = {}
    for source_name, source_rows in captures:
        grouped = {}
        for row in source_rows:
            item = normalize_row(row)
            prior_source = item.get('source_capture', '')
            prior_parts = []
            for source_part in str(prior_source).split('|') if prior_source else []:
                if os.path.isabs(source_part) or source_part.startswith('\\\\') or re.match(r'^[A-Za-z]:[\\/]', source_part):
                    source_part = 'capture-' + hashlib.sha256(source_part.encode('utf-8')).hexdigest()[:20]
                prior_parts.append(source_part)
            prior_source = '|'.join(prior_parts)
            item['source_capture'] = f'{prior_source}|{source_name}' if prior_source else source_name
            grouped.setdefault(item['month'], []).append(item)
        for month, incoming in grouped.items():
            current = by_month.get(month)
            if current is None:
                by_month[month] = incoming
                continue
            old_ids = [row_identity(x) for x in current]
            new_ids = [row_identity(x) for x in incoming]
            common = 0
            for a, b in zip(old_ids, new_ids):
                if a != b: break
                common += 1
            if common and (common == min(len(old_ids), len(new_ids))):
                # Both captures begin at the same list head. Keep the longer list.
                longer = old_ids if len(old_ids) > len(new_ids) else new_ids
                if common and len(longer) > common and longer[common - 1] == longer[common]:
                    raise ValueError(f'A repeated identical row makes the overlap boundary ambiguous in {month}; refusing to discard a possible genuine trip.')
                for offset in range(common):
                    left, right = current[offset], incoming[offset]
                    source_keys = set(FIELDS) - {'list_sequence', 'capture_date', 'source_xml', 'raw_text', 'month'}
                    diffs = {k: [left.get(k, ''), right.get(k, '')] for k in source_keys if left.get(k, '') != right.get(k, '')}
                    if diffs:
                        if conflict == 'error':
                            raise ValueError(f'Conflicting overlapping entries in {month}: {diffs}. Choose --conflict keep-first or keep-last to resolve explicitly.')
                        existing = json.loads(left.get('merge_conflicts') or '[]')
                        existing.append({'source_capture': source_name, 'differences': diffs, 'resolution': conflict})
                        left['merge_conflicts'] = json.dumps(existing, ensure_ascii=False)
                        if conflict == 'keep-last':
                            for key in diffs: left[key] = right[key]
                    left['source_capture'] = left.get('source_capture', '') + '|' + source_name
                if len(new_ids) > len(old_ids):
                    current.extend(incoming[common:])
                continue
            overlap = overlap_size(current, incoming)
            if overlap:
                try:
                    additions = merge_visible(current, incoming, resume=False)
                except ValueError as e:
                    raise ValueError(f'Cannot safely merge {month}: {e}') from e
                for offset in range(overlap):
                    left, right = current[len(current)-overlap+offset], incoming[offset]
                    source_keys = set(FIELDS) - {'list_sequence', 'capture_date', 'source_xml', 'raw_text', 'month'}
                    diffs = {k: [left.get(k, ''), right.get(k, '')] for k in source_keys if left.get(k, '') != right.get(k, '')}
                    if diffs:
                        if conflict == 'error':
                            raise ValueError(f'Conflicting overlapping entries in {month}: {diffs}. Choose --conflict keep-first or keep-last to resolve explicitly.')
                        left.setdefault('merge_conflicts', '[]')
                        existing = json.loads(left['merge_conflicts'] or '[]')
                        existing.append({'source_capture': source_name, 'differences': diffs, 'resolution': conflict})
                        left['merge_conflicts'] = json.dumps(existing, ensure_ascii=False)
                        if conflict == 'keep-last':
                            for key in diffs: left[key] = right[key]
                    left['source_capture'] = left.get('source_capture', '') + '|' + source_name
                current.extend(additions)
                continue
            raise ValueError(f'Cannot establish safe ordered overlap for {month} between captures. Conflict policies only resolve values within records whose sequence alignment is proven.')
    result = []
    for month in sorted(by_month):
        for sequence, row in enumerate(by_month[month], 1):
            row['list_sequence'] = sequence
            result.append(row)
    return result


def load_capture_directory(folder):
    folder = Path(folder).expanduser().resolve()
    data = capture_data_path(folder)
    state_path = data / 'capture-state.json'
    if state_path.exists():
        state = json.loads(state_path.read_text(encoding='utf-8'))
        if state.get('version') != 1: raise ValueError(f'Unsupported state version in {folder}.')
        rows = [normalize_row(row, i + 1) for i, row in enumerate(state.get('rows', []))]
        checks, manifest = state.get('checks', []), state.get('manifest', {})
    else:
        rows = read_capture_rows(data)
        checks_path = data / 'checks.json'
        checks = json.loads(checks_path.read_text(encoding='utf-8')) if checks_path.exists() else []
        manifest_path = data / 'manifest.json'
        manifest = json.loads(manifest_path.read_text(encoding='utf-8')) if manifest_path.exists() else {}
    return folder, rows, checks, manifest


def read_session_manifest(folder):
    """Read the current checkpoint manifest or its standalone compatibility copy."""
    folder = capture_data_path(folder)
    state_path = folder / 'capture-state.json'
    if state_path.exists():
        state = json.loads(state_path.read_text(encoding='utf-8'))
        if state.get('version') != 1:
            raise ValueError(f'Unsupported state version in {folder}.')
        manifest = state.get('manifest')
    else:
        manifest_path = folder / 'manifest.json'
        manifest = json.loads(manifest_path.read_text(encoding='utf-8')) if manifest_path.exists() else None
    if not isinstance(manifest, dict):
        raise ValueError(f'No valid session manifest was found in {folder}.')
    return manifest


def merge_directories(sources, output, conflict='error'):
    loaded = [load_capture_directory(source) for source in sources]
    capture_ids = [capture_identifier(folder, manifest) for folder, _, _, manifest in loaded]
    merged = merge_capture_sets(list(zip(capture_ids, [rows for _, rows, _, _ in loaded])), conflict)
    out = Path(output).expanduser().resolve()
    if any(out == folder for folder, _, _, _ in loaded):
        raise ValueError('Merge output must be a new directory, not one of the source captures.')
    checks_by_month = {}
    source_month_checks = {}
    source_rows_by_month = {}
    for folder, source_rows, checks, manifest in loaded:
        for row in source_rows:
            source_rows_by_month.setdefault((folder, row['month']), []).append(row)
        for entry in checks:
            month = entry.get('month')
            if not month: continue
            source_month_checks.setdefault(month, []).append(entry)
    rows_by_month = {}
    for row in merged:
        rows_by_month.setdefault(row['month'], []).append(row)
    for month, entries in source_month_checks.items():
        complete = [entry for entry in entries if entry.get('status') == 'complete']
        base = dict(complete[-1] if complete else entries[-1])
        for field, label in (('expected_trip_count','trip totals'),('expected_distance_displayed','BMW monthly distances')):
            values = [entry.get(field) for entry in entries if entry.get(field) not in (None,'')]
            distinct = list(dict.fromkeys(values))
            if len(distinct) > 1:
                if conflict == 'error': raise ValueError(f'Conflicting app {label} across captures for {month}: {distinct}.')
                base.setdefault('warnings', []).append(f'Conflicting source {label}: {distinct}; policy={conflict}.')
                base[field] = distinct[0] if conflict == 'keep-first' else distinct[-1]
        base['source_statuses'] = [entry.get('status', 'unknown') for entry in entries]
        checks_by_month[month] = base
    selection = sorted(set(checks_by_month) | {row['month'] for row in merged} |
                       {month for _, _, _, manifest in loaded for month in manifest.get('selection', [])})
    for month in selection:
        checks_by_month.setdefault(month, {'month': month, 'status': 'pending', 'warnings': []})
        check = checks_by_month[month]
        month_rows = rows_by_month.get(month, [])
        represented = sum(int(row['grouped_trip_count']) for row in month_rows)
        expected = check.get('expected_trip_count')
        expected = int(expected) if expected not in (None, '') else None
        check['exported_list_entries'] = len(month_rows)
        check['represented_trip_count'] = represented
        check['distance_sum_km'] = round(sum(parse_distance(row.get('distance_km')) or 0 for row in month_rows), 3)
        if expected is not None and represented != expected:
            check['status'] = 'incomplete'
        elif month in source_month_checks:
            statuses = [entry.get('status', 'unknown') for entry in source_month_checks[month]]
            complete_source_proves_coverage = False
            if expected is not None:
                for folder, _, source_checks, _ in loaded:
                    for source_check in source_checks:
                        if source_check.get('month') != month or source_check.get('status') != 'complete':
                            continue
                        try:
                            source_expected = int(source_check.get('expected_trip_count'))
                        except (TypeError, ValueError):
                            continue
                        source_rows = source_rows_by_month.get((folder, month), [])
                        source_represented = sum(int(row['grouped_trip_count']) for row in source_rows)
                        if source_expected == expected and source_represented == expected:
                            complete_source_proves_coverage = True
                            break
                    if complete_source_proves_coverage:
                        break
            if complete_source_proves_coverage:
                check['status'] = 'complete'
                if any(status != 'complete' for status in statuses):
                    check.setdefault('warnings', []).append('A complete source capture proves month coverage; overlapping partial captures were merged by verified sequence alignment.')
            elif any(status in ('incomplete', 'in_progress') for status in statuses) or month_rows:
                check['status'] = 'incomplete'
                check.setdefault('warnings', []).append('No individually complete source capture proves coverage; matching merged counts alone are insufficient.')
            elif statuses and all(status == 'failed' for status in statuses):
                check['status'] = 'failed'
            else:
                check['status'] = 'pending'
        elif month_rows:
            check['status'] = 'incomplete'
    sessions = [session for _, _, _, manifest in loaded for session in manifest.get('sessions', [])]
    now = dt.datetime.now().astimezone()
    manifest = {'version':1,'kind':'merged','selection':selection,'created_at':now.isoformat(),'timezone':timezone_label(now),
                'timezone_offset':now.strftime('%z'),'sessions':sessions,'source_captures':capture_ids,
                'settings':{'merge_conflict_policy':conflict}}
    out.mkdir(parents=True, exist_ok=False)
    data = out / CAPTURE_DATA_DIR
    (data / 'raw').mkdir(parents=True)
    atomic_json(data / '.owner.json', {'kind':'bmw-journey-exporter-capture-data','created_at':now.isoformat()})
    atomic_json(data / 'capture-state.json', {'version':1,'rows':merged,'checks':[checks_by_month[m] for m in sorted(checks_by_month)],'manifest':manifest})
    atomic_csv(data / 'capture.csv', merged, FIELDS + PROVENANCE_FIELDS)
    atomic_json(data / 'checks.json', [checks_by_month[m] for m in sorted(checks_by_month)])
    write_manifest(data,manifest)
    return out


def capture_identifier(folder, manifest):
    session_ids = sorted(str(session.get('session_id')) for session in manifest.get('sessions', [])
                         if isinstance(session, dict) and session.get('session_id'))
    identity = '\0'.join(session_ids) if session_ids else str(Path(folder).resolve())
    digest = hashlib.sha256(identity.encode('utf-8')).hexdigest()[:20]
    return 'capture-' + digest


def completion_issues(selection, checks, rows):
    """Return requested months that are not proven complete from saved rows."""
    by_month = {str(check.get('month')): check for check in checks if check.get('month')}
    rows_by_month = {}
    for row in rows:
        rows_by_month.setdefault(str(row.get('month', '')), []).append(row)
    issues = []
    for month in selection:
        check = by_month.get(month, {})
        represented = sum(int(row.get('grouped_trip_count') or 1) for row in rows_by_month.get(month, []))
        try:
            expected = int(check['expected_trip_count']) if check.get('expected_trip_count') not in (None, '') else None
        except (TypeError, ValueError):
            expected = None
        if check.get('status') != 'complete' or expected is None or represented != expected:
            issues.append({'month':month, 'status':check.get('status','missing'),
                           'represented':represented, 'expected':expected})
    return issues


def relative_date(text, week, captured_on, month_context=''):
    """Resolve relative dates only when capture date and heading agree."""
    text = clean(text)
    capture = None
    if captured_on:
        try: capture = dt.date.fromisoformat(captured_on)
        except (TypeError, ValueError): return '', 'unresolved: invalid capture date'
    selected_month = None
    if month_context:
        try: selected_month = dt.date.fromisoformat(month_context + '-01')
        except (TypeError, ValueError): selected_month = None
    day_names = {'MON': 0, 'MONDAY': 0, 'TUE': 1, 'TUES': 1, 'TUESDAY': 1, 'WED': 2, 'WEDS': 2,
                 'WEDNESDAY': 2, 'THU': 3, 'THUR': 3, 'THURS': 3, 'THURSDAY': 3, 'FRI': 4,
                 'FRIDAY': 4, 'SAT': 5, 'SATURDAY': 5, 'SUN': 6, 'SUNDAY': 6}
    token = text.split(',', 1)[0].strip().upper().replace('.', '')
    token = re.sub(r'\s+.*$', '', token)
    target_weekday = day_names.get(token)

    def period_from_heading(heading):
        heading = clean(heading).upper()
        if heading == 'THIS WEEK':
            if capture is None: return None
            start = capture - dt.timedelta(days=capture.weekday())
            return start, start + dt.timedelta(days=6)
        if heading == 'LAST WEEK':
            if capture is None: return None
            end = capture - dt.timedelta(days=capture.weekday() + 1)
            return end - dt.timedelta(days=6), end
        m = re.fullmatch(r'WEEK OF\s+(\d{1,2})[./-](\d{1,2})\s*[-–]\s*(\d{1,2})[./-](\d{1,2})(?:[./-](20\d{2}))?', heading)
        if not m:
            return None
        d1, m1, d2, m2 = map(int, m.groups()[:4])
        explicit_year = int(m[5]) if m[5] else None
        candidates = []
        reference_year = selected_month.year if selected_month else (capture.year if capture else None)
        years = [explicit_year] if explicit_year else ([reference_year - 1, reference_year, reference_year + 1] if reference_year else [])
        for year in years:
            try:
                start = dt.date(year, m1, d1)
                end_year = year + (1 if (m2, d2) < (m1, d1) else 0)
                end = dt.date(end_year, m2, d2)
            except ValueError:
                continue
            if selected_month:
                month_end = (selected_month.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
                if start <= month_end and end >= selected_month: candidates.append((start, end))
            elif capture and start <= capture + dt.timedelta(days=1) and end >= capture - dt.timedelta(days=14):
                candidates.append((start, end))
        return candidates[0] if len(candidates) == 1 else None

    period = period_from_heading(week)
    explicit = re.search(r'\b(\d{1,2})[./-](\d{1,2})[./-](20\d{2})\b', text)
    if explicit:
        day, month, year = map(int, explicit.groups())
        try: resolved = dt.date(year, month, day)
        except ValueError: return '', 'unresolved: invalid displayed date'
        if ((capture and resolved > capture) or
                (selected_month and (resolved.year, resolved.month) != (selected_month.year, selected_month.month)) or
                (period and not period[0] <= resolved <= period[1])):
            return '', 'unresolved: displayed date conflicts with capture/week heading'
        if target_weekday is not None and resolved.weekday() != target_weekday:
            return '', 'unresolved: weekday conflicts with displayed date'
        return resolved.isoformat(), 'displayed absolute date'

    if capture is None and not (week and week.upper().startswith('WEEK OF') and selected_month):
        return '', 'unresolved: capture date unknown'

    relative = None
    if re.fullmatch(r'TODAY(?:,.*)?', text, re.I):
        relative = capture
    elif re.fullmatch(r'YESTERDAY(?:,.*)?', text, re.I):
        relative = capture - dt.timedelta(days=1) if capture else None
    elif target_weekday is not None and period is not None:
        matches = [period[0] + dt.timedelta(days=i) for i in range((period[1] - period[0]).days + 1)
                   if (period[0] + dt.timedelta(days=i)).weekday() == target_weekday]
        matches = [d for d in matches if capture is None or d <= capture]
        if len(matches) == 1: relative = matches[0]
    if relative is not None:
        if (capture and relative > capture) or (period and not period[0] <= relative <= period[1]):
            return '', 'unresolved: relative date conflicts with week heading'
        return relative.isoformat(), 'derived: ' + ('capture date + ' if capture else 'week heading ') + (week or 'relative label')
    return '', 'unresolved: source text preserved or week heading ambiguous'


def overlap_size(existing, visible):
    # Sequence overlap, not a global set: identical trips on different dates survive.
    a = [clean(x['raw_text']) for x in existing]
    b = [clean(x['raw_text']) for x in visible]
    for size in range(min(len(a), len(b)), 0, -1):
        if a[-size:] == b[:size]:
            return size
    return 0


def screen_signature(screen):
    return hashlib.sha256(json.dumps([clean(n.get('content-desc', '')) for n in screen['nodes']]).encode()).hexdigest()


def merge_visible(existing, visible, resume=False):
    """Align a page to saved sequence; reject ambiguous duplicate boundaries."""
    if not existing:
        return list(visible)
    size = overlap_size(existing, visible)
    if size:
        old = [clean(x['raw_text']) for x in existing]
        new = [clean(x['raw_text']) for x in visible]
        # If the overlap endpoint consists of identical adjacent rows, content
        # alone cannot tell whether the next row is old or a real repeated trip.
        repeated_inside = (size > 1 and (len(set(old[-size:])) < size or len(set(new[:size])) < size))
        if (repeated_inside or (size < len(old) and old[-size-1] == old[-size]) or
                (size < len(new) and new[size-1] == new[size])):
            raise ValueError('Repeated identical rows make this screen boundary ambiguous. Saved progress is intact; with --keep-debug, inspect capture-data/raw before continuing.')
        return list(visible[size:])
    if resume and visible:
        raws = [clean(r['raw_text']) for r in existing]
        seen = [clean(r['raw_text']) for r in visible]
        if any(raws[i:i + len(seen)] == seen for i in range(len(raws))):
            return []
    if visible:
        raise ValueError('No safe sequence overlap with saved entries; stopping to avoid skipping or duplicating trips.')
    return []


class Exporter:
    def __init__(self, args):
        self.args = args
        configured_timezone = getattr(args, 'timezone', None)
        now = capture_time(timezone_name=configured_timezone)
        self.timezone_name = configured_timezone
        stamp = now.strftime('%Y%m%d-%H%M%S-%f')
        self.out = Path(args.output or ('bmw-journeys-' + stamp)).expanduser().resolve()
        if args.resume and not self.out.is_dir():
            raise ValueError('--resume requires an existing capture directory.')
        self.out.mkdir(parents=True, exist_ok=args.resume)
        nested_data = self.out / CAPTURE_DATA_DIR
        if args.resume:
            self.data_dir = capture_data_path(self.out)
        else:
            nested_data.mkdir()
            self.data_dir = nested_data
        (self.data_dir / 'raw').mkdir(exist_ok=True)
        self.adb = ['adb'] + (['-s', args.serial] if args.serial else [])
        self.rows, self.reports = [], []
        self.months = list(getattr(args, 'months', []))
        self.selection = list(self.months)
        self.manifest_path = self.data_dir / 'manifest.json'
        self.manifest = {
            'version': 1, 'selection': self.selection,
            'created_at': now.isoformat(), 'timezone': timezone_label(now),
            'timezone_offset': now.strftime('%z'), 'sessions': [],
            'settings': {'serial': getattr(args, 'serial', None), 'pause': getattr(args, 'pause', 1.5),
                         'max_pages': getattr(args, 'max_pages', 250), 'keep_awake': getattr(args, 'keep_awake', False),
                         'timezone': configured_timezone,
                         'formats': list(getattr(args, 'formats', ['xlsx'])),
                         'keep_capture': bool(getattr(args, 'keep_capture', False)),
                         'keep_debug': bool(getattr(args, 'keep_debug', False))},
        }
        session = {'session_id': now.strftime('%Y%m%dT%H%M%S%f%z'), 'started_at': now.isoformat(),
                   'timezone': timezone_label(now), 'timezone_offset': now.strftime('%z'),
                   'capture_date': now.date().isoformat(), 'finished_at': None}
        self.manifest['sessions'].append(session)
        self.session = session
        if not (self.data_dir / '.owner.json').exists():
            atomic_json(self.data_dir / '.owner.json', {'kind':'bmw-journey-exporter-capture-data','session_id':session['session_id'],'owned_raw_files':[]})
        saved = self.data_dir / 'capture.csv'
        if args.resume:
            state_path = self.data_dir / 'capture-state.json'
            if state_path.exists():
                state = json.loads(state_path.read_text(encoding='utf-8'))
                if state.get('version') != 1: raise ValueError('Unsupported capture-state version.')
                self.rows = [normalize_row(row, i + 1) for i, row in enumerate(state.get('rows', []))]
                self.reports = state.get('checks', [])
                previous = state.get('manifest', {})
            else:
                if saved.exists() or (self.data_dir / 'journeys.csv').exists():
                    self.rows = read_capture_rows(self.data_dir)
                checks = self.data_dir / 'checks.json'
                if checks.exists(): self.reports = json.loads(checks.read_text(encoding='utf-8'))
                previous = json.loads(self.manifest_path.read_text(encoding='utf-8')) if self.manifest_path.exists() else {}
            if previous:
                if previous.get('version') != 1:
                    raise ValueError(f"Unsupported session manifest version: {previous.get('version')!r}.")
                if previous.get('selection') != self.selection:
                    raise ValueError('Resume selection does not match the saved session manifest.')
                expected = previous.get('settings', {})
                for setting in ('serial', 'pause', 'max_pages', 'keep_awake', 'timezone'):
                    current = self.manifest['settings'][setting]
                    if setting in expected and expected.get(setting) != current:
                        raise ValueError(f'Resume setting --{setting.replace("_", "-")} differs from manifest ({expected.get(setting)!r} vs {current!r}).')
                self.manifest = previous
                self.manifest.setdefault('sessions', []).append(session)
                saved_settings = self.manifest.setdefault('settings', {})
                saved_settings.update(formats=list(getattr(args, 'formats', ['xlsx'])),
                                      keep_capture=bool(getattr(args, 'keep_capture', False)),
                                      keep_debug=bool(getattr(args, 'keep_debug', False)))
                self.selection = previous['selection']
            else:
                # Legacy captures are resumable, but their original run metadata is unknown.
                self.manifest['legacy_capture_metadata'] = 'unknown; source row capture dates were preserved when available'
        existing_ids = []
        for path in (self.data_dir / 'raw').glob('screen-*.xml'):
            m = re.fullmatch(r'screen-(\d+)\.xml', path.name)
            if m: existing_ids.append(int(m[1]))
        self.screen_number = max(existing_ids, default=0)
        self.remote = '/sdcard/bmw-export-' + session['session_id'] + '.xml'
        self.remote_dump_attempted = False
        self.capture_date = now.date().isoformat()
        self.timezone = timezone_label(now)

    def command(self, *parts):
        gesture = len(parts) >= 3 and parts[0] == 'shell' and parts[1] == 'input' and parts[2] in ('tap', 'swipe')
        last = None
        attempts = 1 if gesture else 3
        for attempt in range(attempts):
            try:
                r = subprocess.run(self.adb + list(parts), capture_output=True, timeout=45, env=adb_server_env())
                if r.returncode == 0:
                    return r.stdout
                last = r.stderr.decode(errors='replace').strip() or 'ADB command failed'
            except (subprocess.TimeoutExpired, OSError) as e:
                last = str(e)
            if attempt + 1 < attempts:
                time.sleep(1 + attempt)
        if gesture:
            raise RuntimeError(f'ADB gesture result is ambiguous: {last}. The tap/swipe was not repeated. Check the current screen before resuming.')
        raise RuntimeError(f'ADB failed after {attempts} attempts: {last}. Check USB, authorization, and that the phone is awake.')

    def capture(self):
        last_error = None
        for attempt in range(3):
            self.screen_number += 1
            rel = f'raw/screen-{self.screen_number:05d}.xml'
            try:
                self.remote_dump_attempted = True
                self.command('shell', 'uiautomator', 'dump', self.remote)
                xml = self.command('exec-out', 'cat', self.remote)
                start = xml.find(b'<?xml')
                if start < 0:
                    start = xml.find(b'<hierarchy')
                if start < 0:
                    raise RuntimeError('ADB returned no XML')
                xml = xml[start:]
                atomic_bytes(self.data_dir / rel, xml)
                owner_path = self.data_dir / '.owner.json'
                if owner_path.exists():
                    ownership = json.loads(owner_path.read_text(encoding='utf-8'))
                    owned = ownership.setdefault('owned_raw_files', [])
                    if rel not in owned: owned.append(rel)
                    atomic_json(owner_path, ownership)
                source = f'{CAPTURE_DATA_DIR}/{rel}' if self.data_dir.name == CAPTURE_DATA_DIR else rel
                return parse_screen(xml, source)
            except (RuntimeError, ET.ParseError) as e:
                last_error = e
                time.sleep(1)
        raise RuntimeError(f'Cannot read list screen: {last_error}')

    def swipe(self, screen, down=False):
        # Stay below the month header and above the persistent bottom navigation.
        bottom = min((bounds(n)[1] for n in screen['nodes']
                      if clean(n.get('content-desc', '')) == 'Vehicle tab'), default=2022)
        width = max(bounds(n)[2] for n in screen['nodes'])
        x = int(width * .5)
        low, high = int(bottom * .87), max(screen['month_bounds'][3] + 80, int(bottom * .46))
        y1, y2 = (high, low) if down else (low, high)
        self.command('shell', 'input', 'swipe', str(x), str(y1), str(x), str(y2), '650')
        time.sleep(self.args.pause)

    def checkpoint(self):
        sorted_rows = []
        for month in sorted({r['month'] for r in self.rows}):
            for i, row in enumerate((r for r in self.rows if r['month'] == month), 1):
                sorted_rows.append(dict(row, list_sequence=i))
        state = {'version': 1, 'rows': sorted_rows, 'checks': self.reports, 'manifest': self.manifest}
        atomic_json(self.data_dir / 'capture-state.json', state)
        atomic_csv(self.data_dir / 'capture.csv', sorted_rows, FIELDS + PROVENANCE_FIELDS)
        atomic_json(self.data_dir / 'checks.json', self.reports)
        # New captures keep every recovery artifact in capture-data. For legacy
        # resumes data_dir is the capture root, preserving their existing layout.
        write_manifest(self.data_dir, self.manifest)

    def enable_keep_awake(self):
        value = self.command('shell', 'settings', 'get', 'system', 'screen_off_timeout').decode(errors='replace').strip()
        if not value.isdigit():
            raise RuntimeError(f'Cannot read the existing screen timeout ({value!r}); keep-awake was not changed.')
        sessions = self.manifest.get('sessions', [])
        pending = [record for record in sessions
                   if record is not self.session and record.get('keep_awake_restore_pending')]
        self.inherited_keep_awake_sessions = []
        original_timeout = value
        if pending:
            devices = {record.get('keep_awake_device_serial') for record in pending}
            targets = {str(record.get('keep_awake_original_timeout', '')) for record in pending}
            applied = {str(record.get('keep_awake_applied_timeout', '7200000')) for record in pending}
            restore_guidance = '; '.join(
                manual_timeout_restore_command(record.get('keep_awake_device_serial'), record.get('keep_awake_original_timeout'))
                for record in pending if record.get('keep_awake_device_serial') and record.get('keep_awake_original_timeout'))
            if devices != {self.args.serial} or len(targets) != 1 or '' in targets or len(applied) != 1 or '' in applied:
                raise RuntimeError('Unresolved keep-awake restoration records conflict or belong to another device. No timeout was changed. Review the saved recovery manifest and restore only the matching device/value manually: ' + (restore_guidance or 'no safe command can be inferred.'))
            target = next(iter(targets))
            applied_value = next(iter(applied))
            if value == target:
                # Keep inherited records pending until this session's explicit
                # restoration command succeeds and can checkpoint the result.
                original_timeout = target
                self.inherited_keep_awake_sessions = pending
            elif value == applied_value:
                original_timeout = target
                self.inherited_keep_awake_sessions = pending
            else:
                command = manual_timeout_restore_command(self.args.serial, target)
                raise RuntimeError(f'Unresolved keep-awake restoration exists, but the current timeout is {value} rather than the saved applied value {applied_value}. No timeout was changed. Inspect the phone setting and, if restoring the previous value is appropriate, run manually: {command}')
        self.previous_screen_timeout = original_timeout
        self.session.update(keep_awake_restore_pending=True,
                            keep_awake_original_timeout=original_timeout,
                            keep_awake_applied_timeout='7200000',
                            keep_awake_device_serial=self.args.serial,
                            keep_awake_change_started_at=dt.datetime.now().astimezone().isoformat())
        self.checkpoint()
        if value != '7200000':
            self.command('shell', 'settings', 'put', 'system', 'screen_off_timeout', '7200000')

    def restore_keep_awake(self):
        previous = getattr(self, 'previous_screen_timeout', None)
        if previous is not None:
            self.command('shell', 'settings', 'put', 'system', 'screen_off_timeout', previous)
            self.previous_screen_timeout = None
            self.session['keep_awake_restored_at'] = dt.datetime.now().astimezone().isoformat()
            self.session['keep_awake_restore_pending'] = False
            for record in getattr(self, 'inherited_keep_awake_sessions', []):
                record['keep_awake_restore_pending'] = False
                record['keep_awake_restored_at'] = self.session['keep_awake_restored_at']
                record['keep_awake_restore_note'] = 'restored successfully during resumed session'
            self.inherited_keep_awake_sessions = []

    def cleanup_remote_xml(self):
        if getattr(self, 'remote_dump_attempted', False):
            best_effort_remote_xml_cleanup(self.adb, self.remote, adb_server_env())
            self.remote_dump_attempted = False

    def export_month(self, screen):
        month = screen['month']
        expected = screen['expected_count']
        if expected is None:
            # Ensure the list really starts at the top; bounded attempts.
            for _ in range(12):
                before = screen_signature(screen)
                self.swipe(screen, down=True)
                screen = self.capture()
                if screen['month'] != month:
                    raise RuntimeError('Month unexpectedly changed during scroll')
                if screen_signature(screen) == before:
                    raise RuntimeError('Screen did not change after swipe; stopping without repeating the gesture. Check whether the phone is locked or the list is already at its top.')
                expected = screen['expected_count']
                if expected is not None:
                    break
        if expected is None:
            raise RuntimeError('Cannot find monthly Total trips. Move to the top and restart.')
        report = next((r for r in self.reports if r.get('month') == month), None)
        if report is None:
            report = dict(month=month)
            self.reports.append(report)
        report.update(expected_trip_count=expected,
                      expected_distance_displayed=screen['expected_km'],
                      status='in_progress', warnings=[])
        local = [r for r in self.rows if r['month'] == month] if self.args.resume else []
        stable, previous_signature = 0, None
        print(f'\n{month}: BMW reports {expected} trips', flush=True)
        for page in range(self.args.max_pages):
            if screen['month'] != month:
                raise RuntimeError('Month changed unexpectedly; export stopped.')
            visible = screen['rows']
            try:
                new = merge_visible(local, visible, self.args.resume)
            except ValueError as e:
                raise RuntimeError(str(e)) from e
            for row in new:
                row['capture_date'] = self.capture_date
                row['source_capture'] = self.out.name
                row['source_session_id'] = self.session['session_id']
                if not row['week_context'] and local:
                    row['week_context'] = local[-1]['week_context']
                local.append(row)
                if row not in self.rows:
                    self.rows.append(row)
            signature = screen_signature(screen)
            stable = stable + 1 if signature == previous_signature else 0
            previous_signature = signature
            counted = sum(int(r['grouped_trip_count']) for r in local)
            report.update(exported_list_entries=len(local), represented_trip_count=counted,
                          distance_sum_km=round(sum(parse_distance(r['distance_km']) or 0 for r in local), 3))
            self.checkpoint()
            print(f'  screen {page+1}: {len(local)} entries / {counted} represented trips', flush=True)
            if stable >= 1:
                report['status'] = 'complete' if counted == expected else 'incomplete'
                if counted != expected:
                    report['warnings'].append('Bottom detected but trip count differs from BMW total.')
                report['warnings'].append('Distances are displayed rounded values. Dates/times are preserved verbatim; no start/end locations or times inferred.')
                self.checkpoint()
                return screen
            self.swipe(screen)
            screen = self.capture()
        raise RuntimeError('Maximum page limit reached; partial export preserved.')

    def previous_month(self, screen):
        year, month = map(int, screen['month'].split('-'))
        target = f'{year if month > 1 else year-1}-{month-1 if month > 1 else 12:02d}'
        expected_number = int(target[-2:])
        candidates = [n for n in screen['nodes']
                      if clean(n.get('content-desc', '')) in MONTHS
                      and MONTHS[clean(n.get('content-desc', ''))] == expected_number
                      and n.get('clickable') == 'true']
        if not candidates:
            raise RuntimeError(f'Previous month button unavailable. Cannot reach {target}.')
        x1, y1, x2, y2 = bounds(candidates[0])
        # Flutter exposes a full-width bound for SEPT; its actual button is left.
        x = min(x1 + 35, x2 - 1)
        self.command('shell', 'input', 'tap', str(x), str((y1+y2)//2))
        for _ in range(4):
            time.sleep(self.args.pause)
            result = self.capture()
            if result['month'] == target:
                return result
            if result['month'] != screen['month']:
                raise RuntimeError('Unexpected month selected; stopping.')
        raise RuntimeError(f'Month did not change to {target}. No further taps issued.')

    def run(self):
        screen = self.capture()
        if self.args.months:
            targets = self.args.months
        else:
            targets = months_for_year(self.args.year)
        if screen['month'] not in targets:
            raise RuntimeError(f'Open a month in the requested range ({targets[0]} through {targets[-1]}).')
        for target in targets:
            if not any(r.get('month') == target for r in self.reports):
                self.reports.append(dict(month=target, status='pending', warnings=[]))
        # Navigate backward from the open month, then forward by app navigation is not
        # reliable across app versions. Capture the open month and older requested months.
        index = targets.index(screen['month'])
        todo = list(reversed(targets[:index + 1]))
        if index + 1 < len(targets):
            raise RuntimeError('For a range extending after the open month, start at its latest month; capture navigates backward.')
        for target in todo:
            if screen['month'] != target:
                screen = self.previous_month(screen)
            self.export_month(screen)
            if target != targets[0]:
                screen = self.previous_month(screen)
        print('\nFinished. Completeness status and warnings are included in the XLSX summary.')


def parse_distance(value):
    if value in ('', None): return None
    return float(str(value).replace(' ', '').replace(',', '.'))


def parse_duration(value):
    text = clean(str(value or '')).lower()
    m = re.fullmatch(r'(\d{1,3}):(\d{2})(?::(\d{2}))?', text)
    if m:
        hours, minutes, seconds = int(m[1]), int(m[2]), int(m[3] or 0)
        return dt.timedelta(hours=hours, minutes=minutes, seconds=seconds)
    hours = re.search(r'(\d+)\s*(?:h|hr|hour)', text)
    minutes = re.search(r'(\d+)\s*(?:m|min|minute)', text)
    seconds = re.search(r'(\d+)\s*(?:s|sec|second)', text)
    if any((hours, minutes, seconds)):
        return dt.timedelta(hours=int(hours[1]) if hours else 0, minutes=int(minutes[1]) if minutes else 0, seconds=int(seconds[1]) if seconds else 0)
    return None


def safe_json_value(value):
    if isinstance(value, (dt.date, dt.datetime, dt.timedelta)): return str(value)
    return value


def append_literal_row(worksheet, values):
    """Append values while ensuring caller-provided strings can never become formulas."""
    worksheet.append(values)
    for cell in worksheet[worksheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def load_validated_workbook(path):
    import openpyxl
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad: raise ValueError(f'Corrupt XLSX member: {bad}')
        names = set(archive.namelist())
        required = {'[Content_Types].xml', '_rels/.rels', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', 'xl/styles.xml'}
        if not required.issubset(names): raise ValueError('XLSX is missing required Open XML workbook parts.')
        rel_root = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        relationships = {r.attrib.get('Id'): r.attrib for r in rel_root}
        for rel in relationships.values():
            target = rel.get('Target', '')
            member = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            if member not in names: raise ValueError(f'XLSX relationship target is missing: {target}')
        workbook_xml = ET.fromstring(archive.read('xl/workbook.xml'))
        for sheet in workbook_xml.findall('.//{*}sheet'):
            rel_id = sheet.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            if rel_id not in relationships: raise ValueError(f"XLSX sheet has no relationship: {sheet.attrib.get('name')}")
            rel = relationships[rel_id]
            if 'worksheet' not in rel.get('Type',''): raise ValueError(f"XLSX sheet relationship has wrong type: {sheet.attrib.get('name')}")
            target = rel['Target']
            member = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            sheet_xml = ET.fromstring(archive.read(member))
            if sheet_xml.tag.rsplit('}',1)[-1] != 'worksheet': raise ValueError(f'Invalid worksheet XML: {member}')
        style_rel = next((r for r in relationships.values() if 'styles' in r.get('Type','')), None)
        if not style_rel: raise ValueError('XLSX styles relationship is missing.')
        style_target = style_rel['Target']
        style_member = style_target.lstrip('/') if style_target.startswith('/') else 'xl/' + style_target
        style_root = ET.fromstring(archive.read(style_member))
        if style_root.tag.rsplit('}',1)[-1] != 'styleSheet' or style_root.find('.//{*}cellXfs') is None:
            raise ValueError('XLSX style sheet is malformed.')
    workbook = openpyxl.load_workbook(path, data_only=False)
    if not workbook.sheetnames or workbook.sheetnames[0] != 'Monthly summary':
        raise ValueError('XLSX Monthly summary must be the first sheet.')
    for ws in workbook.worksheets:
        if ws.max_row and ws['A1'].style_id == 0: raise ValueError(f'XLSX header has no style: {ws.title}')
        if ws.max_row > 1 and ws.freeze_panes != 'A2': raise ValueError(f'XLSX header is not frozen: {ws.title}')
    return workbook


def build_reports(folder, replacements=None, formats=None):
    """Build staged reports and replace report files only after all validate."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError('XLSX export requires openpyxl; install dependencies with: python -m pip install -r requirements.txt') from exc
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    folder = Path(folder)
    data = capture_data_path(folder)
    formats = list(formats if formats is not None else EXPORT_FORMATS)
    if not formats or set(formats) - set(EXPORT_FORMATS):
        raise ValueError('Unsupported or empty report format selection.')
    state_path = data / 'capture-state.json'
    if state_path.exists():
        state = json.loads(state_path.read_text(encoding='utf-8'))
        rows = [normalize_row(r, i + 1) for i, r in enumerate(state.get('rows', []))]
        checks = state.get('checks', [])
        manifest = state.get('manifest', {})
    else:
        rows = read_capture_rows(data)
        checks_path = data / 'checks.json'
        checks = json.loads(checks_path.read_text(encoding='utf-8')) if checks_path.exists() else []
        manifest_path = data / 'manifest.json'
        manifest = json.loads(manifest_path.read_text(encoding='utf-8')) if manifest_path.exists() else {}
    normalized_checks = []
    for raw_check in checks:
        if not isinstance(raw_check, dict): raise ValueError('Malformed saved monthly checks; expected objects in checks.json.')
        check = dict(raw_check)
        if check.get('status') == 'in_progress': check['status'] = 'incomplete'
        normalized_checks.append(check)
    checks = normalized_checks
    checks_by_month = {x.get('month'): x for x in checks if x.get('month')}
    checks = [checks_by_month[month] for month in sorted(checks_by_month)]
    replacements = replacements or {}
    normalized = []
    for row in rows:
        item = dict(row)
        resolved, provenance = relative_date(item.get('displayed_date_time', ''), item.get('week_context', ''), item.get('capture_date', ''), item.get('month', ''))
        item['resolved_date'] = resolved
        item['date_basis'] = provenance
        item['destination_label'] = replacements.get(item.get('destination', ''), item.get('destination', ''))
        item['start_location_label'] = replacements.get(item.get('start_location', ''), item.get('start_location', ''))
        item['distance_km_numeric'] = parse_distance(item.get('distance_km'))
        item['duration_numeric'] = parse_duration(item.get('duration_displayed'))
        normalized.append(item)
    months = sorted(set(manifest.get('selection', [])) | set(checks_by_month) | {r['month'] for r in normalized})
    if not months: months = sorted({r['month'] for r in normalized})
    report_fields = list(dict.fromkeys([*FIELDS, *PROVENANCE_FIELDS, 'resolved_date', 'date_basis', 'destination_label', 'start_location_label']))
    serialized = [{k: safe_json_value(v) for k, v in row.items()} for row in normalized]

    with tempfile.TemporaryDirectory(prefix='.reports-', dir=folder) as staging_name:
        staging = Path(staging_name)
        csv_fields = report_fields
        atomic_csv(staging / 'journeys.csv', normalized, csv_fields)
        atomic_json(staging / 'journeys.json', {'entries': serialized, 'monthly_checks': checks, 'requested_months': months})
        xml_root = ET.Element('bmwJourneyExport')
        for row in serialized:
            entry = ET.SubElement(xml_root, 'entry')
            for key in report_fields: ET.SubElement(entry, key).text = str(row.get(key, '') or '')
        ET.indent(xml_root, space='  ')
        ET.ElementTree(xml_root).write(staging / 'journeys.xml', encoding='utf-8', xml_declaration=True)

        wb = openpyxl.Workbook()
        summary = wb.active; summary.title = 'Monthly summary'
        summary_headers = ['Month', 'Status', 'Warnings', 'Captured entries', 'Represented trips', 'BMW trip count', 'BMW distance (km)', 'BMW distance source', 'Displayed entry distance sum (km)', 'Distance comparison (km)']
        summary.append(summary_headers)
        for month in months:
            subset = [r for r in normalized if r['month'] == month]
            check = checks_by_month.get(month, {})
            bmw_distance = parse_distance(check.get('expected_distance_displayed'))
            entry_sum = round(sum(r['distance_km_numeric'] or 0 for r in subset), 3)
            represented = sum(int(r['grouped_trip_count']) for r in subset)
            status = check.get('status', 'pending' if not subset else 'incomplete')
            expected_count = check.get('expected_trip_count')
            expected_count = int(expected_count) if expected_count not in (None, '') else None
            warning_parts = [str(w) for w in check.get('warnings', []) if w]
            conflict_resolutions = []
            for row in subset:
                try:
                    conflict_resolutions.extend(json.loads(row.get('merge_conflicts') or '[]'))
                except (TypeError, ValueError):
                    continue
            if conflict_resolutions:
                counts = {}
                for conflict_record in conflict_resolutions:
                    resolution = conflict_record.get('resolution', 'explicit policy') if isinstance(conflict_record, dict) else 'explicit policy'
                    counts[resolution] = counts.get(resolution, 0) + 1
                resolved_text = ', '.join(f'{count} resolved with {policy}' for policy, count in sorted(counts.items()))
                warning_parts.append(f'Overlapping-entry conflicts: {resolved_text}.')
            warnings = '; '.join(warning_parts)
            append_literal_row(summary, [month, status, warnings, len(subset), represented, expected_count, bmw_distance, check.get('expected_distance_displayed', ''), entry_sum, (round(entry_sum - bmw_distance, 3) if bmw_distance is not None else None)])
        if months:
            r = summary.max_row + 1
            summary.cell(r, 1, 'TOTAL')
            for col in (4,5,6,7,9,10): summary.cell(r,col, f'=SUM({get_column_letter(col)}2:{get_column_letter(col)}{r-1})')
        detail_headers = ['month','list_sequence','week_context','destination','destination_label','displayed_date_time','resolved_date','date_basis','distance_km_numeric','distance_km_source','duration_numeric','duration_displayed','consumption_displayed','grouped_trip_count','start_location','departure_time','arrival_time']
        header_fill = PatternFill('solid', fgColor='174A6E')
        group_fill = PatternFill('solid', fgColor='FFF2CC')
        total_fill = PatternFill('solid', fgColor='D9EAD3')
        for cell in summary[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=header_fill; cell.alignment=Alignment(wrap_text=True)
        summary.freeze_panes='A2'; summary.auto_filter.ref=f'A1:{get_column_letter(summary.max_column)}{summary.max_row-1}'
        for month in months:
            ws=wb.create_sheet(month)
            ws.append(detail_headers)
            for row in [r for r in normalized if r['month']==month]:
                resolved=dt.date.fromisoformat(row['resolved_date']) if row['resolved_date'] else None
                duration=row['duration_numeric']
                append_literal_row(ws, [row['month'],row['list_sequence'],row['week_context'],row['destination'],row['destination_label'],row['displayed_date_time'],resolved,row['date_basis'],row['distance_km_numeric'],row['distance_km'],duration,row['duration_displayed'],row['consumption_displayed'],int(row['grouped_trip_count']),row['start_location'],row['departure_time'],row['arrival_time']])
                if int(row['grouped_trip_count']) > 1:
                    for cell in ws[ws.max_row]: cell.fill=group_fill
            for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=header_fill; cell.alignment=Alignment(wrap_text=True)
            last_data=ws.max_row
            ws.append(['MONTH TOTAL'] + ['']*7 + [f'=SUM(I2:I{last_data})' if last_data > 1 else 0, '', f'=SUM(K2:K{last_data})' if last_data > 1 else 0, '', '', f'=SUM(N2:N{last_data})' if last_data > 1 else 0])
            for cell in ws[ws.max_row]: cell.fill=total_fill; cell.font=Font(bold=True)
            ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(detail_headers))}{max(1,last_data)}'
            for row_index in range(2, last_data + 1):
                ws.cell(row_index, 7).number_format='yyyy-mm-dd'
                ws.cell(row_index, 9).number_format='0.000'
                ws.cell(row_index, 11).number_format='[h]:mm:ss'
            for column, number_format in ((7,'yyyy-mm-dd'),(9,'0.000'),(11,'[h]:mm:ss')):
                ws.cell(last_data + 1, column).number_format=number_format
            for col, header in enumerate(detail_headers,1):
                values=[len(str(header))]+[len(str(ws.cell(r,col).value or '')) for r in range(2, min(ws.max_row,202)+1)]
                ws.column_dimensions[get_column_letter(col)].width=min(max(max(values)+2,12),48)
        summary_widths=[13,14,48,17,20,17,20,22,28,22]
        for col,width in enumerate(summary_widths,1): summary.column_dimensions[get_column_letter(col)].width=width
        for r in range(2,summary.max_row+1):
            for c in (4,5,6): summary.cell(r,c).number_format='0'
            for c in (7,9,10): summary.cell(r,c).number_format='#,##0.000'
        wb.save(staging/'journeys.xlsx')
        load_validated_workbook(staging/'journeys.xlsx').close()

        if 'csv' in formats:
            with (staging/'journeys.csv').open(encoding='utf-8-sig', newline='') as stream:
                if len(list(csv.DictReader(stream))) != len(normalized): raise ValueError('CSV report entry count validation failed.')
        if 'json' in formats:
            if len(json.loads((staging/'journeys.json').read_text(encoding='utf-8'))['entries']) != len(normalized):
                raise ValueError('JSON report entry count validation failed.')
        if 'xml' in formats:
            if len(ET.parse(staging/'journeys.xml').getroot().findall('entry')) != len(normalized):
                raise ValueError('XML report entry count validation failed.')
        for fmt in formats:
            if not (staging / {'csv':'journeys.csv','json':'journeys.json','xml':'journeys.xml','xlsx':'journeys.xlsx'}[fmt]).is_file():
                raise ValueError(f'{fmt.upper()} report was not saved.')

        names_by_format={'csv':'journeys.csv','json':'journeys.json','xml':'journeys.xml','xlsx':'journeys.xlsx'}
        names=[names_by_format[fmt] for fmt in formats]
        backup=staging/'backup'; backup.mkdir()
        old=[]; installed=[]
        try:
            for name in names:
                target=folder/name
                if target.exists():
                    saved=backup/name; os.replace(target,saved); old.append((target,saved))
            for name in names:
                os.replace(staging/name,folder/name); installed.append(folder/name)
        except Exception:
            for target in installed:
                if target.exists(): target.unlink()
            for target,saved in old:
                if saved.exists(): os.replace(saved,target)
            raise
    return normalized


def doctor(serial=None):
    print('ADB executable:', shutil.which('adb') or 'not found')
    try:
        rows = adb_devices(serial)
    except RuntimeError as e:
        print('ADB check:', e)
        return 1
    host, port = adb_server_address()
    print(f'ADB server: reachable at {host}:{port}')
    if not rows:
        print('Devices: none found' + (f' for serial {serial}' if serial else ''))
        return 1
    ok = True
    for row in rows:
        device, state = row[0], row[1] if len(row) > 1 else 'unknown'
        print(f'Device {device}: {state}')
        if state != 'device':
            print('  Check USB debugging and accept the authorization prompt on the phone.')
            ok = False
            continue
        try:
            screen = screen_for_device(device)
            print(f"  My BMW English All journeys screen: yes ({screen['month']})")
        except Exception as e:
            print(f'  My BMW English All journeys screen: no ({e})')
            ok = False
    return 0 if ok else 1


def main():
    configure_console_streams()
    p = argparse.ArgumentParser(description='Capture the English My BMW journey list through an already running ADB server, then generate reports.', allow_abbrev=False)
    p.add_argument('--doctor', action='store_true', help='Check ADB server, device authorization, and the current list screen without navigating.')
    p.add_argument('--report-only', metavar='DIR', help='Regenerate reports from saved data without connecting a phone.')
    p.add_argument('--formats', type=parse_formats, default=None, metavar='SELECTION', help='Reports to generate: xlsx (default), comma-separated formats, or all (xlsx,csv,json,xml).')
    p.add_argument('--merge', action='append', default=[], metavar='DIR', help='Merge a capture directory; repeat for multiple sources, with --output.')
    p.add_argument('--conflict', choices=('error','keep-first','keep-last'), default='error', help='Explicit policy for conflicting capture records.')
    p.add_argument('--year', type=int, help='Capture Jan through Dec; current year stops at the current month.')
    p.add_argument('--month', help='Capture one month, YYYY-MM.')
    p.add_argument('--from-month', help='Inclusive range start, YYYY-MM.')
    p.add_argument('--to-month', help='Inclusive range end, YYYY-MM.')
    p.add_argument('--serial', help='ADB device serial when several are connected.')
    p.add_argument('--timezone', metavar='IANA_ZONE', help='Capture timezone, for example Europe/Ljubljana (default: host local timezone).')
    p.add_argument('--output', help='Capture directory (new unless --resume).')
    p.add_argument('--resume', action='store_true', help='Resume after validating the manifest selection and settings.')
    p.add_argument('--keep-capture', action=argparse.BooleanOptionalAction, default=None, help='Retain checkpoint, CSV and manifest data after a complete export.')
    p.add_argument('--keep-debug', action=argparse.BooleanOptionalAction, default=None, help='Retain capture data and raw screen XML; implies --keep-capture.')
    p.add_argument('--keep-awake', action=argparse.BooleanOptionalAction, default=None, help='Temporarily extend Android screen timeout and restore the previous value at exit.')
    p.add_argument('--pause', type=float)
    p.add_argument('--max-pages', type=int)
    p.add_argument('--replace-label', action='append', default=[], metavar='OLD=NEW')
    args = p.parse_args()
    try:
        replacements = {}
        for value in args.replace_label:
            if '=' not in value or not value.split('=', 1)[0]: p.error('--replace-label must be OLD=NEW')
            k, v = value.split('=', 1); replacements[k] = v
        if args.keep_debug:
            args.keep_capture = True
        if args.doctor:
            if args.report_only or args.merge or args.resume or args.month or args.year is not None or args.from_month or args.to_month:
                p.error('--doctor cannot be combined with capture, report, or merge options.')
            return doctor(args.serial)
        if args.report_only:
            if args.merge or args.output or args.resume or args.year is not None or args.month or args.from_month or args.to_month or args.keep_capture or args.keep_debug:
                p.error('--report-only cannot be combined with capture selection, --merge, --output, or --resume.')
            build_reports(args.report_only, replacements, args.formats or ['xlsx'])
            print(f'Reports generated in {Path(args.report_only).resolve()}')
            return 0
        if args.merge:
            if not args.output: p.error('--output is required with --merge.')
            if args.resume: p.error('--merge creates a new capture; omit --resume.')
            if args.year is not None or args.month or args.from_month or args.to_month or args.serial or args.keep_awake:
                p.error('--merge cannot be combined with phone capture options.')
            if args.keep_debug: args.keep_capture = True
            output = merge_directories(args.merge, args.output, args.conflict)
            try:
                build_reports(output, replacements, args.formats or ['xlsx'])
            except Exception as e:
                print(f'Merge report generation or validation failed; recovery data retained in {Path(output) / CAPTURE_DATA_DIR}: {e}', file=sys.stderr)
                return 1
            _, merged_rows, merged_checks, merged_manifest = load_capture_directory(output)
            issues = completion_issues(merged_manifest.get('selection', []), merged_checks, merged_rows)
            if issues:
                print(f'Merge is incomplete; recovery data retained in {Path(output) / CAPTURE_DATA_DIR}.')
                for issue in issues:
                    expected = 'unknown' if issue['expected'] is None else str(issue['expected'])
                    print(f"  {issue['month']}: status={issue['status']}, represented={issue['represented']}, expected={expected}")
                return 2
            cleanup_capture_data(output, bool(args.keep_capture), bool(args.keep_debug))
            print(f'Merged {len(args.merge)} captures and generated reports in {output}')
            return 0
        selectors = (args.year is not None, args.month is not None, args.from_month is not None or args.to_month is not None)
        if sum(selectors) > 1 or (args.from_month is None) != (args.to_month is None):
            p.error('Choose one of --year, --month, or both --from-month and --to-month.')
        saved_manifest = None
        if args.resume:
            if not args.output: p.error('--output is required with --resume.')
            if not Path(args.output).expanduser().is_dir(): p.error('--resume requires an existing output directory.')
            saved_manifest = read_session_manifest(args.output)
            if saved_manifest.get('kind') == 'merged' or saved_manifest.get('source_captures'):
                p.error('Merged datasets are reportable and mergeable, but cannot resume phone capture. Start a new phone capture with its own selection.')
            if saved_manifest.get('version') != 1: p.error('Unsupported session manifest version.')
            settings = saved_manifest.get('settings', {})
            if not isinstance(settings, dict): p.error('Malformed session settings in manifest.')
            for attribute, key, default in (('serial','serial',None),('timezone','timezone',None),('pause','pause',1.5),('max_pages','max_pages',250),('keep_awake','keep_awake',False)):
                supplied = getattr(args, attribute)
                if supplied is not None and key in settings and supplied != settings[key]:
                    if key == 'serial':
                        pending = [record for record in saved_manifest.get('sessions', [])
                                   if isinstance(record, dict) and record.get('keep_awake_restore_pending')]
                        if pending:
                            guidance = '; '.join(manual_timeout_restore_command(record.get('keep_awake_device_serial'), record.get('keep_awake_original_timeout'))
                                                 for record in pending if record.get('keep_awake_device_serial') and record.get('keep_awake_original_timeout'))
                            p.error('The requested device differs from the device with an unresolved keep-awake restoration. No timeout was changed. Resume with the original device, or inspect the phone and restore only the saved device/value manually: ' + (guidance or 'the saved manifest lacks safe restoration details.'))
                    p.error(f'--{key.replace("_", "-")} conflicts with the saved session setting {settings[key]!r}.')
                if supplied is None:
                    setattr(args, attribute, settings.get(key, default))
            for key, default in (('formats',['xlsx']),('keep_capture',False),('keep_debug',False)):
                supplied = getattr(args, key)
                saved = settings.get(key, default)
                if key == 'formats':
                    try:
                        if isinstance(saved, list) and all(isinstance(item, str) for item in saved):
                            saved = parse_formats(','.join(saved))
                        elif isinstance(saved, str):
                            saved = parse_formats(saved)
                        else:
                            raise argparse.ArgumentTypeError('Malformed saved formats.')
                    except argparse.ArgumentTypeError as exc:
                        p.error(f'Malformed report formats in the saved manifest: {exc}')
                if supplied is not None and key == 'formats' and list(supplied) != list(saved):
                    p.error('--formats conflicts with the saved session selection.')
                if supplied is None: setattr(args, key, list(saved) if key == 'formats' else saved)
            if args.keep_debug: args.keep_capture = True
        else:
            if args.formats is None: args.formats = ['xlsx']
            if args.keep_capture is None: args.keep_capture = False
            if args.keep_debug is None: args.keep_debug = False
            if args.keep_debug: args.keep_capture = True
            if args.pause is None: args.pause = 1.5
            if args.max_pages is None: args.max_pages = 250
            if args.keep_awake is None: args.keep_awake = False
        if args.pause < .5: p.error('--pause must be at least 0.5 seconds')
        if args.max_pages < 1: p.error('--max-pages must be positive')
        try:
            today = capture_time(timezone_name=args.timezone).date()
        except ValueError as exc:
            p.error(str(exc))
        if args.resume and not any(selectors):
            saved_selection = saved_manifest.get('selection')
            if not isinstance(saved_selection, list) or not saved_selection:
                p.error('The saved manifest has no date selection; provide --year, --month, or a complete month range.')
            if month_range(saved_selection[0], saved_selection[-1]) != saved_selection:
                p.error('The saved manifest date selection is malformed or non-contiguous.')
            args.months = saved_selection
        else:
            try:
                if args.month:
                    args.months = month_range(args.month, args.month)
                elif args.from_month:
                    args.months = month_range(args.from_month, args.to_month)
                else:
                    args.year = args.year or today.year
                    args.months = months_for_year(args.year, today)
            except ValueError as exc:
                p.error(str(exc))
        if saved_manifest and args.months != saved_manifest.get('selection'):
            p.error('Explicit resume date selection conflicts with the saved session selection.')
        try:
            validate_requested_months(args.months, today)
        except ValueError as exc:
            p.error(str(exc))
        if not args.output: p.error('--output is required for capture; use --report-only for saved data.')
        # Check the existing server without starting or restarting it.
        devices = adb_devices(args.serial)
        ready = [d[0] for d in devices if len(d) > 1 and d[1] == 'device']
        if args.serial and args.serial not in ready:
            raise RuntimeError(f'Device {args.serial} is unavailable or unauthorized. Unlock the phone and accept its USB debugging prompt.')
        if not args.serial:
            if not ready: raise RuntimeError('No authorized ADB device. Unlock the phone and accept the USB debugging prompt.')
            if len(ready) > 1: raise RuntimeError('Several devices are connected: ' + ', '.join(ready) + '. Rerun with --serial SERIAL.')
            args.serial = ready[0]
        exporter = Exporter(args)
        for month in args.months:
            if not any(r.get('month') == month for r in exporter.reports):
                exporter.reports.append({'month': month, 'status': 'pending', 'warnings': []})
        exporter.checkpoint()
        print('Output:', exporter.out, flush=True)
        print('Open the latest requested month in My BMW → All journeys. The UI must be English. Keep the phone unlocked; Ctrl+C saves progress.', flush=True)
        result = 0
        previous_sigterm = signal.getsignal(signal.SIGTERM)
        def stop_on_sigterm(signum, frame):
            raise KeyboardInterrupt('Received SIGTERM; saving recovery data.')
        signal.signal(signal.SIGTERM, stop_on_sigterm)
        try:
            if args.keep_awake: exporter.enable_keep_awake()
            exporter.run()
        except (Exception, KeyboardInterrupt) as e:
            print(f'\nSTOPPED: {e or "Interrupted"}. Recovery data remains in {exporter.data_dir}.', file=sys.stderr)
            for report in exporter.reports:
                if report.get('month') and report.get('status') in ('pending', 'in_progress'):
                    report['status'] = 'incomplete' if any(r['month'] == report['month'] for r in exporter.rows) else 'failed'
                    report.setdefault('warnings', []).append(str(e))
            exporter.reports.append({'status': 'failed' if not exporter.rows else 'incomplete', 'reason': str(e)})
            result = 1
        finally:
            cleanup_remote = getattr(exporter, 'cleanup_remote_xml', None)
            if cleanup_remote: cleanup_remote()
            try: exporter.restore_keep_awake()
            except Exception as e:
                print(f'Could not restore phone screen timeout: {e}', file=sys.stderr)
                exporter.session['keep_awake_restore_error'] = str(e)
                original_timeout = exporter.session.get('keep_awake_original_timeout')
                serial = exporter.session.get('keep_awake_device_serial')
                if original_timeout and serial:
                    print('Manual restore command:', manual_timeout_restore_command(serial, original_timeout), file=sys.stderr)
                for report in exporter.reports:
                    if report.get('month'): report.setdefault('warnings', []).append('Phone screen timeout restoration failed: ' + str(e))
                result = 1
            exporter.session['finished_at'] = dt.datetime.now().astimezone().isoformat()
            try:
                exporter.checkpoint()
            finally:
                signal.signal(signal.SIGTERM, previous_sigterm)
        try:
            build_reports(exporter.out, replacements, args.formats)
        except Exception as e:
            print(f'Capture is saved in {exporter.data_dir}, but report generation or validation failed safely: {e}', file=sys.stderr)
            retry_path = host_path_for_recovery(args.output)
            quoted_output = shell_quote(retry_path)
            quoted_report_path = shell_quote(host_path_for_recovery(args.output, 'report'))
            print(f"Report retry: make report INPUT={quoted_report_path} FORMATS={','.join(args.formats)}", file=sys.stderr)
            print(f"Report retry (local Python): python3 bmw_export_journeys.py --report-only {quoted_output} --formats {','.join(args.formats)}", file=sys.stderr)
            for command in resume_commands(retry_path): print('Resume command:', command, file=sys.stderr)
            return 1
        issues = completion_issues(args.months, exporter.reports, exporter.rows)
        complete = not issues
        if result or not complete:
            print('Incomplete export; recovery data retained.')
            for issue in issues:
                expected = 'unknown' if issue['expected'] is None else str(issue['expected'])
                print(f"  {issue['month']}: status={issue['status']}, represented={issue['represented']}, expected={expected}")
            for command in resume_commands(host_path_for_recovery(args.output)):
                print('Resume command:', command)
            return result or 2
        cleanup_capture_data(exporter.out, args.keep_capture, args.keep_debug)
        if args.keep_debug:
            print(f'Complete export. Capture data and raw XML retained in {exporter.data_dir}.')
        elif args.keep_capture:
            print(f'Complete export. Capture data retained in {exporter.data_dir}; raw XML removed.')
        else:
            print('Complete export. Temporary capture data removed.')
        return 0
    except (ValueError, RuntimeError) as e:
        print(f'Error: {e}', file=sys.stderr)
        output_value = getattr(args, 'output', None)
        if output_value:
            recovery = Path(output_value).expanduser() / CAPTURE_DATA_DIR
            if recovery.is_dir(): print(f'Recovery data remains in {recovery}.', file=sys.stderr)
        return 1
    except OSError as e:
        print(f'Filesystem error: {e}', file=sys.stderr)
        output_value = getattr(args, 'output', None)
        if output_value:
            recovery = Path(output_value).expanduser() / CAPTURE_DATA_DIR
            if recovery.is_dir(): print(f'Recovery data remains in {recovery}.', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
