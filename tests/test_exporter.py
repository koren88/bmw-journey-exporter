import csv
import json
import argparse
import contextlib
import datetime as dt
import importlib.util
import io
import os
import signal
import tempfile
import subprocess
import sys
import unittest
from unittest import mock
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import bmw_export_journeys as app

FIXTURE = Path(__file__).parent / 'fixtures' / 'sanitized_screen.xml'
HAS_OPENPYXL = importlib.util.find_spec('openpyxl') is not None

class ExporterTests(unittest.TestCase):
    def make_merge_sources(self, root, expected=1, status='complete'):
        row=app.parse_screen(FIXTURE.read_bytes())['rows'][0]
        row.update(month='2026-10',list_sequence=1,capture_date='2026-10-21',grouped_trip_count=1,raw_text='Fictional journey')
        sources=[]
        for index in range(2):
            source=root/f'merge-source-{index}'; source.mkdir(); sources.append(source)
            app.atomic_csv(source/'capture.csv',[dict(row)],app.FIELDS+app.PROVENANCE_FIELDS)
            app.write_manifest(source,{'selection':['2026-10'],'settings':{},'sessions':[]})
            app.atomic_json(source/'checks.json',[{'month':'2026-10','status':status,'expected_trip_count':expected}])
        return sources

    def run_merge_cli(self, sources, output, extra=(), report_error=None):
        argv=['bmw_export_journeys.py','--merge',str(sources[0]),'--merge',str(sources[1]),'--output',str(output),*extra]
        with mock.patch.object(app.sys,'argv',argv), \
             mock.patch.object(app,'build_reports',side_effect=report_error) if report_error else mock.patch.object(app,'build_reports'):
            return app.main()

    def test_merge_cli_incomplete_retains_recovery_and_returns_two(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); sources=self.make_merge_sources(root,expected=2,status='incomplete'); out=root/'out'
            with contextlib.redirect_stdout(io.StringIO()) as stdout:
                result=self.run_merge_cli(sources,out)
            self.assertEqual(result,2)
            self.assertIn('Merge is incomplete',stdout.getvalue())
            self.assertTrue((out/'capture-data'/'capture-state.json').exists())

    def test_merge_cli_success_cleans_recovery_and_keep_capture_retains_it(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); sources=self.make_merge_sources(root); out=root/'clean'
            with contextlib.redirect_stdout(io.StringIO()):
                result=self.run_merge_cli(sources,out)
            self.assertEqual(result,0)
            self.assertFalse((out/'capture-data').exists())
            kept=root/'kept'
            with contextlib.redirect_stdout(io.StringIO()):
                result=self.run_merge_cli(sources,kept,['--keep-capture'])
            self.assertEqual(result,0)
            self.assertTrue((kept/'capture-data'/'capture-state.json').exists())

    def test_merge_cli_report_failure_retains_recovery(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); sources=self.make_merge_sources(root); out=root/'report-fails'
            with contextlib.redirect_stderr(io.StringIO()) as stderr:
                result=self.run_merge_cli(sources,out,report_error=OSError('simulated workbook failure'))
            self.assertEqual(result,1)
            self.assertIn('recovery data retained',stderr.getvalue())
            self.assertTrue((out/'capture-data'/'capture-state.json').exists())

    def test_confirmed_resume_initialization_regressions(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); (folder/'raw').mkdir()
            (folder/'raw'/'screen-00001.xml').write_text('old', encoding='utf-8')
            app.atomic_csv(folder/'capture.csv',[],app.FIELDS)
            args=argparse.Namespace(output=str(folder),resume=True,serial=None)
            exporter=app.Exporter(args)
            self.assertTrue(exporter.screen_number >= 1)
            self.assertTrue((folder/'raw'/'screen-00001.xml').exists())

    def test_interrupted_resume_orchestration_uses_mocked_adb_and_preserves_raw(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); (folder/'raw').mkdir()
            old_raw=folder/'raw'/'screen-00001.xml'; old_raw.write_text('preserve-me', encoding='utf-8')
            parsed=app.parse_screen(FIXTURE.read_bytes())
            old_row=parsed['rows'][0]
            old_row.update(list_sequence=1,capture_date='2025-01-08',source_capture=folder.name,source_session_id='old')
            manifest={'version':1,'selection':['2025-01'],'created_at':'2025-01-08T08:00:00+01:00','timezone':'CET','timezone_offset':'+0100','settings':{'serial':'MOCK','pause':0.5,'max_pages':5,'keep_awake':False},'sessions':[]}
            state={'version':1,'rows':[old_row],'checks':[{'month':'2025-01','status':'incomplete','warnings':[]}],'manifest':manifest}
            app.atomic_json(folder/'capture-state.json',state)
            args=argparse.Namespace(output=str(folder),resume=True,serial='MOCK',months=['2025-01'],pause=.5,max_pages=5,year=2025,keep_awake=False)
            exporter=app.Exporter(args)
            xml=FIXTURE.read_bytes()
            def fake_run(argv,**kwargs):
                if 'input' in argv: return subprocess.CompletedProcess(argv,0,b'',b'')
                if 'uiautomator' in argv: return subprocess.CompletedProcess(argv,0,b'',b'')
                if argv[-2:] == ['cat', exporter.remote]: return subprocess.CompletedProcess(argv,0,xml,b'')
                return subprocess.CompletedProcess(argv,0,b'',b'')
            with mock.patch.object(app.subprocess,'run',side_effect=fake_run):
                exporter.run()
            self.assertEqual(old_raw.read_text(encoding='utf-8'),'preserve-me')
            self.assertTrue((folder/'raw'/'screen-00002.xml').exists())
            self.assertEqual(sum(int(r['grouped_trip_count']) for r in exporter.rows),3)
            self.assertEqual(next(r for r in exporter.reports if r.get('month')=='2025-01')['status'],'complete')

    def test_ambiguous_adb_gesture_is_never_retried(self):
        exporter=app.Exporter.__new__(app.Exporter)
        exporter.adb=['adb','-s','MOCK']
        result=subprocess.CompletedProcess(['adb'],1,b'',b'connection reset')
        with mock.patch.object(app.subprocess,'run',return_value=result) as run:
            with self.assertRaisesRegex(RuntimeError,'gesture result is ambiguous'):
                exporter.command('shell','input','tap','100','200')
        self.assertEqual(run.call_count,1)

    def test_adb_preflight_does_not_start_server_if_socket_is_absent(self):
        with mock.patch.object(app.shutil,'which',return_value='/fake/adb'), \
             mock.patch.object(app.socket,'create_connection',side_effect=OSError('not listening')), \
             mock.patch.object(app.subprocess,'run') as run:
            with self.assertRaisesRegex(RuntimeError,'never starts or restarts'):
                app.adb_devices()
        run.assert_not_called()

    def test_keep_awake_persists_restore_details_before_setting_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            args=argparse.Namespace(output=str(Path(tmp)/'capture'),resume=False,serial='MOCK-SERIAL',months=['2026-10'],timezone='UTC',pause=1,max_pages=5,keep_awake=True,formats=['xlsx'],keep_capture=False,keep_debug=False)
            exporter=app.Exporter(args)
            calls=[]
            def fake_command(*parts):
                calls.append(parts)
                if parts == ('shell','settings','get','system','screen_off_timeout'):
                    return b'60000'
                state=json.loads((exporter.data_dir/'capture-state.json').read_text(encoding='utf-8'))
                session=state['manifest']['sessions'][-1]
                self.assertEqual(session['keep_awake_original_timeout'],'60000')
                self.assertEqual(session['keep_awake_device_serial'],'MOCK-SERIAL')
                self.assertTrue(session['keep_awake_restore_pending'])
                return b''
            with mock.patch.object(exporter,'command',side_effect=fake_command):
                exporter.enable_keep_awake()
            self.assertEqual(calls[-1],('shell','settings','put','system','screen_off_timeout','7200000'))
            self.assertEqual(app.manual_timeout_restore_command('MOCK-SERIAL','60000'),"adb -s 'MOCK-SERIAL' shell settings put system screen_off_timeout '60000'")
            with mock.patch.object(exporter,'command',side_effect=RuntimeError('device disconnected')):
                with self.assertRaisesRegex(RuntimeError,'disconnected'):
                    exporter.restore_keep_awake()
            self.assertTrue(exporter.session['keep_awake_restore_pending'])

    def test_keep_awake_pending_timeout_is_inherited_and_restored_after_resume(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'
            def arguments(resume=False):
                return argparse.Namespace(output=str(out),resume=resume,serial='MOCK-SERIAL',months=['2026-10'],timezone='UTC',pause=1,max_pages=5,keep_awake=True,formats=['xlsx'],keep_capture=True,keep_debug=False)
            timeout=['60000']
            first=app.Exporter(arguments())
            def first_command(*parts):
                if parts == ('shell','settings','get','system','screen_off_timeout'): return timeout[0].encode()
                if parts[-1] == '7200000': timeout[0]='7200000'
                return b''
            with mock.patch.object(first,'command',side_effect=first_command): first.enable_keep_awake()
            self.assertEqual(timeout[0],'7200000')  # hard kill here: no finally restoration
            resumed=app.Exporter(arguments(True))
            calls=[]
            def resumed_command(*parts):
                calls.append(parts)
                if parts == ('shell','settings','get','system','screen_off_timeout'): return timeout[0].encode()
                if parts[-1] == '60000': timeout[0]='60000'
                elif parts[-1] == '7200000': timeout[0]='7200000'
                return b''
            with mock.patch.object(resumed,'command',side_effect=resumed_command):
                resumed.enable_keep_awake()
                self.assertEqual(resumed.session['keep_awake_original_timeout'],'60000')
                self.assertEqual(calls.count(('shell','settings','put','system','screen_off_timeout','7200000')),0)
                resumed.restore_keep_awake()
            self.assertEqual(timeout[0],'60000')
            prior=resumed.manifest['sessions'][0]
            self.assertFalse(prior['keep_awake_restore_pending'])
            self.assertFalse(resumed.session['keep_awake_restore_pending'])

    def test_pending_keep_awake_restoration_failure_preserves_both_records(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'
            def arguments(resume=False):
                return argparse.Namespace(output=str(out),resume=resume,serial='MOCK-SERIAL',months=['2026-10'],timezone='UTC',pause=1,max_pages=5,keep_awake=True,formats=['xlsx'],keep_capture=True,keep_debug=False)
            first=app.Exporter(arguments())
            def first_command(*parts): return b'60000' if parts == ('shell','settings','get','system','screen_off_timeout') else b''
            with mock.patch.object(first,'command',side_effect=first_command): first.enable_keep_awake()
            resumed=app.Exporter(arguments(True))
            def resumed_command(*parts):
                if parts == ('shell','settings','get','system','screen_off_timeout'): return b'7200000'
                raise RuntimeError('mock restore failure')
            with mock.patch.object(resumed,'command',side_effect=resumed_command):
                resumed.enable_keep_awake()
                with self.assertRaisesRegex(RuntimeError,'restore failure'): resumed.restore_keep_awake()
            self.assertTrue(resumed.manifest['sessions'][0]['keep_awake_restore_pending'])
            self.assertTrue(resumed.session['keep_awake_restore_pending'])
            resumed.checkpoint()

    def test_pending_keep_awake_changed_timeout_and_different_device_are_explicit(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'
            def arguments(serial='MOCK-SERIAL',resume=False):
                return argparse.Namespace(output=str(out),resume=resume,serial=serial,months=['2026-10'],timezone='UTC',pause=1,max_pages=5,keep_awake=True,formats=['xlsx'],keep_capture=True,keep_debug=False)
            first=app.Exporter(arguments())
            def first_command(*parts): return b'60000' if parts == ('shell','settings','get','system','screen_off_timeout') else b''
            with mock.patch.object(first,'command',side_effect=first_command): first.enable_keep_awake()
            resumed=app.Exporter(arguments(resume=True))
            calls=[]
            def changed_command(*parts):
                calls.append(parts)
                return b'120000' if parts == ('shell','settings','get','system','screen_off_timeout') else b''
            with mock.patch.object(resumed,'command',side_effect=changed_command):
                with self.assertRaisesRegex(RuntimeError,'current timeout is 120000') as failure: resumed.enable_keep_awake()
            self.assertIn("adb -s 'MOCK-SERIAL'",str(failure.exception))
            self.assertEqual(calls,[('shell','settings','get','system','screen_off_timeout')])
            stderr=io.StringIO()
            with mock.patch.object(app.sys,'argv',['cli','--resume','--output',str(out),'--serial','OTHER']), \
                 contextlib.redirect_stderr(stderr):
                with self.assertRaises(SystemExit) as exit_error: app.main()
            self.assertEqual(exit_error.exception.code,2)
            self.assertIn('unresolved keep-awake restoration',stderr.getvalue())
            self.assertIn("adb -s 'MOCK-SERIAL'",stderr.getvalue())

    def test_sigterm_during_capture_checkpoints_and_reports_interruption(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'
            def terminate(_exporter):
                signal.raise_signal(signal.SIGTERM)
            def fake_reports(folder,replacements,formats):
                (Path(folder)/'journeys.xlsx').write_text('mocked report', encoding='utf-8')
            with mock.patch.object(app.sys,'argv',['cli','--month','2026-10','--output',str(out)]), \
                 mock.patch.object(app,'adb_devices',return_value=[['MOCK','device']]), \
                 mock.patch.object(app.Exporter,'run',terminate), \
                 mock.patch.object(app,'build_reports',side_effect=fake_reports), \
                 contextlib.redirect_stdout(io.StringIO()) as stdout, contextlib.redirect_stderr(io.StringIO()) as stderr:
                result=app.main()
            self.assertEqual(result,1)
            self.assertIn('Received SIGTERM',stderr.getvalue())
            self.assertTrue((out/'capture-data'/'capture-state.json').exists())

    def test_restore_failure_prints_saved_manual_adb_command(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'
            def adb_command(_exporter,*parts):
                if parts == ('shell','settings','get','system','screen_off_timeout'): return b'60000'
                if parts[-1] == '60000': raise RuntimeError('device disconnected')
                return b''
            def fail_capture(_exporter): raise RuntimeError('simulated screen failure')
            def fake_reports(folder,replacements,formats): (Path(folder)/'journeys.xlsx').write_text('mocked report', encoding='utf-8')
            with mock.patch.object(app.sys,'argv',['cli','--month','2026-10','--output',str(out),'--keep-awake']), \
                 mock.patch.object(app,'adb_devices',return_value=[['MOCK-SERIAL','device']]), \
                 mock.patch.object(app.Exporter,'command',adb_command), \
                 mock.patch.object(app.Exporter,'run',fail_capture), \
                 mock.patch.object(app,'build_reports',side_effect=fake_reports), \
                 contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()) as stderr:
                self.assertEqual(app.main(),1)
            self.assertIn("adb -s 'MOCK-SERIAL' shell settings put system screen_off_timeout '60000'",stderr.getvalue())
            state=json.loads((out/'capture-data'/'capture-state.json').read_text(encoding='utf-8'))
            self.assertTrue(state['manifest']['sessions'][-1]['keep_awake_restore_pending'])

    def test_doctor_removes_only_its_temporary_xml_even_after_parse_failure(self):
        calls=[]
        def fake_run(command,**kwargs):
            calls.append(command)
            if 'dump' in command: return subprocess.CompletedProcess(command,0,b'',b'')
            if 'cat' in command: return subprocess.CompletedProcess(command,0,b'<not-xml/>',b'')
            return subprocess.CompletedProcess(command,1,b'',b'device offline')
        with mock.patch.object(app.subprocess,'run',side_effect=fake_run), contextlib.redirect_stderr(io.StringIO()) as stderr:
            with self.assertRaisesRegex(RuntimeError,'no UI XML'):
                app.screen_for_device('MOCK')
        self.assertIn('device offline',stderr.getvalue())
        remote=calls[0][-1]
        self.assertEqual(calls[-1],['adb','-s','MOCK','shell','rm','-f',remote])
        self.assertEqual(remote,f'/sdcard/bmw-doctor-{os.getpid()}.xml')

    def test_saved_rows_are_normalized_and_malformed_data_rejected(self):
        row={'month':'2026-10','list_sequence':'2','grouped_trip_count':'3','distance_km':'12.5','capture_date':'2026-10-21','raw_text':'fiction'}
        normalized=app.normalize_row(row, 2)
        self.assertEqual(normalized['grouped_trip_count'],3)
        self.assertEqual(normalized['list_sequence'],2)
        with self.assertRaises(ValueError):
            app.normalize_row(dict(row, grouped_trip_count='three'), 2)

    def test_resume_manifest_selection_and_settings_are_validated(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); (folder/'raw').mkdir()
            app.atomic_json(folder/'capture-state.json', {'version':1,'rows':[],'checks':[],'manifest':{'version':1,'selection':['2025-01'],'settings':{'serial':'A','pause':1.5,'max_pages':250,'keep_awake':False},'sessions':[]}})
            args=argparse.Namespace(output=str(folder),resume=True,serial='A',months=['2025-02'],pause=1.5,max_pages=250,keep_awake=False)
            with self.assertRaisesRegex(ValueError,'selection'):
                app.Exporter(args)
            args.months=['2025-01']; args.pause=2
            with self.assertRaisesRegex(ValueError,'pause'):
                app.Exporter(args)

    def test_resume_timezone_must_match_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); (folder/'raw').mkdir()
            app.atomic_json(folder/'capture-state.json', {'version':1,'rows':[],'checks':[],'manifest':{'version':1,'selection':['2025-01'],'settings':{'serial':'A','pause':1.5,'max_pages':250,'keep_awake':False,'timezone':'Europe/Ljubljana'},'sessions':[]}})
            args=argparse.Namespace(output=str(folder),resume=True,serial='A',months=['2025-01'],pause=1.5,max_pages=250,keep_awake=False,timezone='UTC')
            with self.assertRaisesRegex(ValueError,'timezone'):
                app.Exporter(args)

    def test_cli_resume_loads_selection_and_settings_from_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)/'capture'; folder.mkdir(); (folder/'raw').mkdir()
            selection=['2026-09','2026-10']
            settings={'serial':'MOCK','pause':0.75,'max_pages':42,'keep_awake':True,'timezone':'Europe/Ljubljana'}
            app.atomic_json(folder/'capture-state.json', {'version':1,'rows':[],'checks':[],'manifest':{'version':1,'selection':selection,'settings':settings,'sessions':[]}})
            observed={}
            class FakeExporter:
                def __init__(self,args):
                    observed.update(vars(args)); self.out=folder; self.reports=[]; self.rows=[]; self.session={}
                def checkpoint(self): pass
                def run(self): pass
                def enable_keep_awake(self): pass
                def restore_keep_awake(self): pass
            with mock.patch.object(app.sys,'argv',['bmw_export_journeys.py','--resume','--output',str(folder)]), \
                 mock.patch.object(app,'adb_devices',return_value=[['MOCK','device']]), \
                 mock.patch.object(app,'Exporter',FakeExporter), \
                 mock.patch.object(app,'build_reports'):
                result=app.main()
            self.assertEqual(result,2)
            self.assertEqual(observed['months'],selection)
            self.assertEqual(observed['serial'],'MOCK')
            self.assertEqual(observed['pause'],0.75)
            self.assertEqual(observed['max_pages'],42)
            self.assertTrue(observed['keep_awake'])
            self.assertEqual(observed['timezone'],'Europe/Ljubljana')

    def test_current_year_selection_and_future_validation(self):
        self.assertEqual(app.months_for_year(2026, dt.date(2026,10,21))[-1], '2026-10')
        with self.assertRaisesRegex(ValueError, 'Future month'):
            app.validate_requested_months(['2026-11'], dt.date(2026,10,21))

    def test_weekday_names_explicit_weeks_and_capture_date_unknown(self):
        daynames=['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday','Sunday']
        expected=['2026-10-12','2026-10-13','2026-10-14','2026-10-15','2026-10-16','2026-10-17','2026-10-18']
        for name,date in zip(daynames, expected):
            self.assertEqual(app.relative_date(name+', 09:00','WEEK OF 12/10 - 18/10','2026-10-21')[0], date)
        self.assertEqual(app.relative_date('Friday, 09:00','LAST WEEK','2026-10-21')[0], '2026-10-16')
        self.assertEqual(app.relative_date('Wednesday, 09:00','WEEK OF 12/10 - 18/10','2026-10-21')[0], '2026-10-14')
        for abbreviation, date in [('Tue','2026-10-13'),('Tues','2026-10-13'),('Wed','2026-10-14'),('Thu','2026-10-15'),('Thurs','2026-10-15'),('Sat','2026-10-17')]:
            self.assertEqual(app.relative_date(abbreviation+', 09:00','WEEK OF 12/10 - 18/10','2026-10-21')[0],date)
        self.assertEqual(app.relative_date('Wednesday, 09:00','WEEK OF 28/12 - 03/01','2027-01-05')[0], '2026-12-30')
        self.assertEqual(app.relative_date('TODAY','THIS WEEK','')[0], '')
        self.assertEqual(app.relative_date('Wednesday, 09:00','WEEK OF 12/10 - 18/10','','2026-10')[0], '2026-10-14')
        self.assertEqual(app.relative_date('20.10.2026, 09:00','WEEK OF 12/10 - 18/10','2026-10-21')[1].split(':')[0], 'unresolved')

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_manifest_months_include_empty_and_pending(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            app.write_manifest(folder, {'selection':['2026-09','2026-10'],'created_at':'x','timezone':'UTC','settings':{}})
            (folder/'capture.csv').write_text(','.join(app.FIELDS)+'\n', encoding='utf-8')
            app.atomic_json(folder/'checks.json',[{'month':'2026-09','status':'in_progress','expected_trip_count':0}])
            app.build_reports(folder)
            workbook=app.load_validated_workbook(folder/'journeys.xlsx')
            months=[r[0] for r in workbook['Monthly summary'].iter_rows(min_row=2,values_only=True) if r[0] and r[0] != 'TOTAL']
            self.assertEqual(months,['2026-09','2026-10'])
            self.assertEqual(workbook['Monthly summary']['B2'].value,'incomplete')
            saved_checks=json.loads((folder/'checks.json').read_text(encoding='utf-8'))
            self.assertEqual(saved_checks[0]['status'],'in_progress')
            self.assertIn('Warnings',[cell.value for cell in workbook['Monthly summary'][1]])
            self.assertEqual(workbook['Monthly summary']['B2'].value,'incomplete')

    def test_cross_capture_merge_conflict_policy_and_provenance(self):
        one=[dict(raw_text='A',month='2026-10',destination='One')]
        two=[dict(raw_text='A',month='2026-10',destination='One'),dict(raw_text='B',month='2026-10',destination='Two')]
        merged=app.merge_capture_sets([('fictional-a',one),('fictional-b',two)])
        self.assertEqual([r['raw_text'] for r in merged],['A','B'])
        self.assertIn('fictional-b',merged[0]['source_capture'])
        with self.assertRaisesRegex(ValueError,'conflict'):
            app.merge_capture_sets([('a',[dict(raw_text='same',month='2026-10',destination='Home')]),('b',[dict(raw_text='same',month='2026-10',destination='Work')])])

    def test_unproven_alignment_is_rejected_even_with_keep_policy(self):
        first=[dict(raw_text='A',month='2026-10'),dict(raw_text='B',month='2026-10')]
        second=[dict(raw_text='C',month='2026-10')]
        for policy in ('keep-first','keep-last'):
            with self.subTest(policy=policy), self.assertRaisesRegex(ValueError,'safe ordered overlap'):
                app.merge_capture_sets([('a',first),('b',second)],conflict=policy)

    def test_merged_month_totals_are_recomputed_from_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); sources=[]
            for name, ids in (('one',['A','B']),('two',['B','C'])):
                source=root/name; source.mkdir(); sources.append(source)
                rows=[dict(month='2026-01',list_sequence=i,destination=identity,raw_text=identity,grouped_trip_count=1,distance_km='1',capture_date='2026-01-15') for i,identity in enumerate(ids,1)]
                app.atomic_csv(source/'capture.csv',rows,app.FIELDS)
                app.atomic_json(source/'checks.json',[{'month':'2026-01','status':'complete','expected_trip_count':2,'expected_distance_displayed':'2','represented_trip_count':2,'distance_sum_km':2}])
                app.write_manifest(source,{'version':1,'selection':['2026-01'],'settings':{},'sessions':[]})
            output=app.merge_directories(sources,root/'merged')
            checks=json.loads((output/'capture-data'/'checks.json').read_text(encoding='utf-8'))
            check=checks[0]
            self.assertEqual(sum(int(r['grouped_trip_count']) for r in app.read_capture_rows(output/'capture-data')),3)
            self.assertEqual(check['represented_trip_count'],3)
            self.assertEqual(check['exported_list_entries'],3)
            self.assertEqual(check['distance_sum_km'],3.0)
            self.assertEqual(check['status'],'incomplete')

    def test_merge_complete_source_proves_coverage_with_overlapping_partial_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); sources=[]
            for name, ids, status in (('complete',['A','B'],'complete'),('partial',['B'],'incomplete')):
                source=root/name; source.mkdir(); sources.append(source)
                rows=[dict(month='2026-01',list_sequence=i,destination=item,raw_text=item,grouped_trip_count=1) for i,item in enumerate(ids,1)]
                app.atomic_csv(source/'capture.csv',rows,app.FIELDS)
                app.atomic_json(source/'checks.json',[{'month':'2026-01','status':status,'expected_trip_count':2}])
                app.write_manifest(source,{'selection':['2026-01'],'settings':{},'sessions':[]})
            merged=app.merge_directories(sources,root/'merged')
            check=json.loads((merged/'capture-data'/'checks.json').read_text(encoding='utf-8'))[0]
            self.assertEqual(check['status'],'complete')
            self.assertEqual(check['represented_trip_count'],2)
            self.assertTrue(any('complete source capture proves' in warning for warning in check['warnings']))

    def test_matching_partial_sources_do_not_fabricate_complete_coverage(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); sources=[]
            for name, ids in (('partial-a',['A']),('partial-b',['A','B'])):
                source=root/name; source.mkdir(); sources.append(source)
                rows=[dict(month='2026-01',list_sequence=i,destination=item,raw_text=item,grouped_trip_count=1) for i,item in enumerate(ids,1)]
                app.atomic_csv(source/'capture.csv',rows,app.FIELDS)
                app.atomic_json(source/'checks.json',[{'month':'2026-01','status':'incomplete','expected_trip_count':2}])
                app.write_manifest(source,{'selection':['2026-01'],'settings':{},'sessions':[]})
            merged=app.merge_directories(sources,root/'merged')
            check=json.loads((merged/'capture-data'/'checks.json').read_text(encoding='utf-8'))[0]
            self.assertEqual(check['represented_trip_count'],2)
            self.assertEqual(check['status'],'incomplete')
            self.assertTrue(any('matching merged counts alone are insufficient' in warning for warning in check['warnings']))

    def test_merge_preserves_noncontiguous_month_selection_and_refuses_resume(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); sources=[]
            for index in range(2):
                source=root/f'source-{index}'; source.mkdir(); sources.append(source)
                rows=[dict(month=m,list_sequence=1,destination=m,raw_text=m,grouped_trip_count=1) for m in ('2026-01','2026-03')]
                app.atomic_csv(source/'capture.csv',rows,app.FIELDS)
                app.atomic_json(source/'checks.json',[{'month':m,'status':'complete','expected_trip_count':1} for m in ('2026-01','2026-03')])
                app.write_manifest(source,{'selection':['2026-01','2026-03'],'settings':{},'sessions':[]})
            merged=app.merge_directories(sources,root/'merged')
            manifest=app.read_session_manifest(merged)
            self.assertEqual(manifest['selection'],['2026-01','2026-03'])
            with mock.patch.object(app.sys,'argv',['cli','--resume','--output',str(merged)]), contextlib.redirect_stderr(io.StringIO()):
                with self.assertRaises(SystemExit) as error: app.main()
            self.assertEqual(error.exception.code,2)

    def test_absolute_date_resolves_without_capture_date(self):
        self.assertEqual(app.relative_date('01/01/2026, 11:32','','','2026-01'),('2026-01-01','displayed absolute date'))
        self.assertTrue(app.relative_date('01/01/2026, 11:32','','','2026-02')[1].startswith('unresolved:'))

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_xlsx_source_strings_are_written_as_literal_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            row=dict(month='2026-01',list_sequence=1,week_context='',destination='=1+1',displayed_date_time='=1+1',distance_km='1',duration_displayed='=1+1',consumption_displayed='',grouped_trip_count=1,start_location='',departure_time='',arrival_time='',raw_text='=1+1',source_xml='raw/screen.xml',capture_date='')
            app.atomic_csv(folder/'capture.csv',[row],app.FIELDS)
            app.atomic_json(folder/'checks.json',[{'month':'2026-01','status':'=1+1','expected_trip_count':1,'expected_distance_displayed':'1'}])
            app.write_manifest(folder,{'version':1,'selection':['2026-01'],'settings':{},'sessions':[]})
            app.build_reports(folder,{'=1+1':'=1+1'},formats=['xlsx','csv','json'])
            wb=app.load_validated_workbook(folder/'journeys.xlsx')
            ws=wb['2026-01']; headers={cell.value:cell.column for cell in ws[1]}
            self.assertNotIn('raw_text',headers)
            for header in ('destination','destination_label','displayed_date_time','duration_displayed'):
                cell=ws.cell(2,headers[header])
                self.assertEqual(cell.value,'=1+1')
                self.assertEqual(cell.data_type,'s')
            self.assertEqual(ws.cell(2,headers['distance_km_source']).data_type,'s')
            summary=wb['Monthly summary']
            self.assertEqual(summary['B2'].value,'=1+1')
            self.assertEqual(summary['B2'].data_type,'s')
            self.assertEqual(summary['H2'].value,'1')
            self.assertEqual(summary['H2'].data_type,'s')
            self.assertEqual(summary['D3'].data_type,'f')
            machine=json.loads((folder/'journeys.json').read_text(encoding='utf-8'))
            self.assertEqual(machine['entries'][0]['raw_text'],'=1+1')
            with (folder/'journeys.csv').open(encoding='utf-8-sig',newline='') as stream:
                csv_row=next(csv.DictReader(stream))
            self.assertEqual(csv_row['raw_text'],'=1+1')
            wb.close()

    def test_explicit_capture_timezone_controls_month_boundary(self):
        instant=dt.datetime(2026,1,1,0,30,tzinfo=dt.timezone.utc)
        local=app.capture_time(instant,'America/Los_Angeles')
        self.assertEqual(local.date().isoformat(),'2025-12-31')
        self.assertEqual(local.tzinfo.key,'America/Los_Angeles')
        self.assertEqual(app.months_for_year(local.year,local.date())[-1],'2025-12')
        with tempfile.TemporaryDirectory() as tmp:
            args=argparse.Namespace(output=str(Path(tmp)/'capture'),resume=False,serial='MOCK',months=['2025-12'],timezone='America/Los_Angeles',pause=1.5,max_pages=250,keep_awake=False)
            with mock.patch.object(app,'capture_time',return_value=local):
                exporter=app.Exporter(args)
            self.assertEqual(exporter.manifest['settings']['timezone'],'America/Los_Angeles')
            self.assertEqual(exporter.session['timezone'],'America/Los_Angeles')
            self.assertEqual(exporter.session['capture_date'],'2025-12-31')

    def test_merge_capture_directories_preserves_session_provenance(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); sources=[]
            rows=[app.parse_screen(FIXTURE.read_bytes())['rows'][0],app.parse_screen(FIXTURE.read_bytes())['rows'][1]]
            for index, subset in enumerate(([rows[0]],[rows[0],rows[1]]),1):
                source=root/f'source-{index}'; source.mkdir(); sources.append(source)
                for i,row in enumerate(subset,1): row.update(list_sequence=i,capture_date='2025-01-08',source_session_id=f'session-{index}',source_capture=str(source))
                app.atomic_csv(source/'capture.csv',subset,app.FIELDS+app.PROVENANCE_FIELDS)
                app.write_manifest(source,{'selection':['2025-01'],'settings':{},'sessions':[{'session_id':f'session-{index}'}]})
                app.atomic_json(source/'checks.json',[{'month':'2025-01','status':'incomplete','expected_trip_count':3}])
            out=app.merge_directories(sources,root/'merged')
            state=json.loads((out/'capture-data'/'capture-state.json').read_text(encoding='utf-8'))
            self.assertEqual([r['raw_text'] for r in state['rows']],[r['raw_text'] for r in rows])
            self.assertIn('·',state['rows'][0]['raw_text'])
            self.assertIn(app.capture_identifier(sources[0],{'sessions':[{'session_id':'session-1'}]}),state['rows'][0]['source_capture'])
            self.assertIn(app.capture_identifier(sources[1],{'sessions':[{'session_id':'session-2'}]}),state['rows'][0]['source_capture'])
            self.assertNotIn(str(root),state['manifest']['source_captures'][0])
            self.assertNotIn(str(root),state['rows'][0]['source_capture'])
            self.assertEqual(len(state['manifest']['sessions']),2)
            cli_out=root/'merged-cli'
            argv=['cli','--merge',str(sources[0]),'--merge',str(sources[1]),'--output',str(cli_out)]
            with mock.patch.object(app.sys,'argv',argv), mock.patch.object(app,'build_reports'), contextlib.redirect_stdout(io.StringIO()) as stdout:
                cli_result=app.main()
            self.assertEqual(cli_result,2)
            self.assertIn('Merge is incomplete',stdout.getvalue())
            self.assertTrue((cli_out/'capture-data'/'capture-state.json').exists())

    def test_main_makes_console_output_safe_for_restrictive_encodings(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'
            stdout_bytes=io.BytesIO(); stderr_bytes=io.BytesIO()
            stdout=io.TextIOWrapper(stdout_bytes,encoding='ascii',errors='strict')
            stderr=io.TextIOWrapper(stderr_bytes,encoding='ascii',errors='strict')
            def finish_capture(exporter):
                exporter.reports[0].update(status='complete',expected_trip_count=0)
            def write_mock_report(folder,replacements,formats):
                (Path(folder)/'journeys.xlsx').write_text('mock report',encoding='utf-8')
            with mock.patch.object(app.sys,'argv',['cli','--month','2026-10','--output',str(out)]), \
                 mock.patch.object(app,'adb_devices',return_value=[['MOCK','device']]), \
                 mock.patch.object(app.Exporter,'run',finish_capture), \
                 mock.patch.object(app,'build_reports',side_effect=write_mock_report), \
                 mock.patch.object(app.sys,'stdout',stdout), mock.patch.object(app.sys,'stderr',stderr):
                self.assertEqual(app.main(),0)
                stdout.flush(); stderr.flush()
                output=stdout_bytes.getvalue()
                self.assertIn(b'\\u2192',output)
                self.assertIn(b'All journeys',output)

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_xlsx_has_numeric_cells_and_valid_relationships(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            row=dict(month='2026-10',list_sequence=1,week_context='THIS WEEK',destination='Office',displayed_date_time='Wednesday, 09:00',distance_km='12.5',duration_displayed='01:30',consumption_displayed='15',grouped_trip_count=2,start_location='',departure_time='',arrival_time='',raw_text='fixture',source_xml='raw/test.xml',capture_date='2026-10-21')
            app.atomic_csv(folder/'capture.csv',[row], app.FIELDS + app.PROVENANCE_FIELDS)
            app.write_manifest(folder, {'months':['2026-10'],'created_at':'x','timezone':'UTC','settings':{}})
            app.atomic_json(folder/'checks.json',[{'month':'2026-10','status':'complete','expected_trip_count':2,'expected_distance_displayed':'12,5'}])
            app.build_reports(folder)
            wb=app.load_validated_workbook(folder/'journeys.xlsx')
            ws=wb['2026-10']
            headers={cell.value:cell.column for cell in ws[1]}
            self.assertIsInstance(ws.cell(2,headers['distance_km_numeric']).value,(int,float))
            self.assertIsInstance(ws.cell(2,headers['duration_numeric']).value,(dt.time,dt.timedelta,int,float))
            self.assertIsInstance(ws.cell(2,headers['resolved_date']).value,dt.date)
            self.assertEqual(ws.cell(2,headers['distance_km_numeric']).number_format,'0.000')
            self.assertEqual(ws.cell(2,headers['duration_numeric']).number_format,'[h]:mm:ss')
            self.assertEqual(ws.cell(2,headers['resolved_date']).number_format,'yyyy-mm-dd')

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_xlsx_detail_columns_are_simple_and_merge_conflicts_are_summary_warnings(self):
        expected=['month','list_sequence','week_context','destination','destination_label','displayed_date_time','resolved_date','date_basis','distance_km_numeric','distance_km_source','duration_numeric','duration_displayed','consumption_displayed','grouped_trip_count','start_location','departure_time','arrival_time']
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); data=root/'capture-data'; data.mkdir()
            row=dict(month='2026-10',list_sequence=1,week_context='THIS WEEK',destination='Fictional place',destination_label='Fictional place',displayed_date_time='Wednesday, 09:00',distance_km='12.5',duration_displayed='01:30',consumption_displayed='15',grouped_trip_count=1,start_location='',departure_time='',arrival_time='',raw_text='sanitized source',source_xml='raw/screen-00001.xml',capture_date='2026-10-21',source_capture='fictional-source',source_session_id='fictional-session',merge_conflicts=json.dumps([{'resolution':'keep-last','differences':{'destination':['A','B']}}]))
            app.atomic_csv(data/'capture.csv',[row],app.FIELDS+app.PROVENANCE_FIELDS)
            app.write_manifest(data,{'selection':['2026-10'],'settings':{}})
            app.atomic_json(data/'checks.json',[{'month':'2026-10','status':'complete','expected_trip_count':1,'warnings':['Existing completeness warning.']}])
            app.build_reports(root)
            wb=app.load_validated_workbook(root/'journeys.xlsx')
            ws=wb['2026-10']
            self.assertEqual([cell.value for cell in ws[1]],expected)
            summary=wb['Monthly summary']
            warnings=summary.cell(2,3).value
            self.assertIn('Existing completeness warning.',warnings)
            self.assertIn('Overlapping-entry conflicts: 1 resolved with keep-last.',warnings)
            saved=app.read_capture_rows(data)[0]
            self.assertEqual(saved['raw_text'],'sanitized source')
            self.assertEqual(saved['source_xml'],'raw/screen-00001.xml')
            self.assertEqual(saved['source_capture'],'fictional-source')
            self.assertEqual(saved['source_session_id'],'fictional-session')
            self.assertTrue(saved['merge_conflicts'])
            wb.close()

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_successful_default_capture_writes_only_xlsx_and_nested_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'successful-capture'
            row=app.parse_screen(FIXTURE.read_bytes())['rows'][0]
            row.update(month='2026-10',list_sequence=1,capture_date='2026-10-21',source_capture='fictional-capture',source_session_id='fictional-session')
            def fake_run(exporter):
                exporter.rows=[dict(row)]
                exporter.reports[0].update(status='complete',expected_trip_count=1)
            with mock.patch.object(app.sys,'argv',['cli','--month','2026-10','--output',str(out)]), \
                 mock.patch.object(app,'adb_devices',return_value=[['MOCK','device']]), \
                 mock.patch.object(app.Exporter,'run',fake_run), \
                 contextlib.redirect_stdout(io.StringIO()):
                result=app.main()
            self.assertEqual(result,0)
            self.assertEqual(sorted(path.name for path in out.iterdir()),['journeys.xlsx'])
            self.assertFalse((out/'manifest.json').exists())
            self.assertFalse((out/'capture-data').exists())

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_report_failure_does_not_damage_checkpoint(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            row=app.parse_screen(FIXTURE.read_bytes())['rows'][0]
            row.update(list_sequence=1,capture_date='2025-01-08')
            app.atomic_csv(folder/'capture.csv',[row],app.FIELDS+app.PROVENANCE_FIELDS)
            state=b'{"version":1,"rows":[],"checks":[],"manifest":{"selection":["2025-01"]}}'
            (folder/'capture-state.json').write_bytes(state)
            from openpyxl import Workbook
            with mock.patch.object(Workbook,'save',side_effect=OSError('simulated report failure')):
                with self.assertRaises(OSError): app.build_reports(folder)
            self.assertEqual((folder/'capture-state.json').read_bytes(),state)

    def test_parse_rows_and_group_count(self):
        parsed = app.parse_screen(FIXTURE.read_bytes())
        self.assertEqual(parsed['month'], '2025-01')
        self.assertEqual([r['grouped_trip_count'] for r in parsed['rows']], [1, 2])
        self.assertEqual(parsed['expected_count'], 3)

    def test_inclusive_month_range(self):
        self.assertEqual(app.month_range('2024-11', '2025-02'), ['2024-11','2024-12','2025-01','2025-02'])
        with self.assertRaises(ValueError): app.month_range('2025-02', '2025-01')

    def test_relative_dates_use_capture_date_and_week(self):
        self.assertEqual(app.relative_date('YESTERDAY', '', '2025-01-08')[0], '2025-01-07')
        self.assertEqual(app.relative_date('MONDAY, 08:15', 'THIS WEEK', '2025-01-08')[0], '2025-01-06')
        self.assertEqual(app.relative_date('MONDAY, 08:15', 'LAST WEEK', '2025-01-08')[0], '2024-12-30')
        self.assertTrue(app.relative_date('TODAY', '', '2025-01-08')[1].startswith('derived:'))

    def test_overlapping_captures_preserve_identical_repeated_rows(self):
        a = [{'raw_text':'same'}, {'raw_text':'same'}, {'raw_text':'other'}]
        b = [{'raw_text':'other'}, {'raw_text':'same'}, {'raw_text':'same'}]
        self.assertEqual(app.overlap_size(a, b), 1)
        self.assertEqual(app.overlap_size(a, [{'raw_text':'same'}]), 0)

    def test_resume_skips_saved_pages_and_appends_new_rows(self):
        saved=[{'raw_text':'A'},{'raw_text':'B'},{'raw_text':'C'}]
        self.assertEqual(app.merge_visible(saved, [{'raw_text':'A'},{'raw_text':'B'}], resume=True), [])
        self.assertEqual(app.merge_visible(saved, [{'raw_text':'C'},{'raw_text':'D'}], resume=True), [{'raw_text':'D'}])

    def test_resume_refuses_ambiguous_repeated_boundary(self):
        saved=[{'raw_text':'A'},{'raw_text':'same'}]
        page=[{'raw_text':'same'},{'raw_text':'same'}]
        with self.assertRaisesRegex(ValueError, 'ambiguous'):
            app.merge_visible(saved, page, resume=True)
        with self.assertRaisesRegex(ValueError, 'ambiguous'):
            app.merge_visible([{'raw_text':'A'},{'raw_text':'same'},{'raw_text':'same'}], page, resume=True)

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_report_totals_and_formats_offline(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            rows = app.parse_screen(FIXTURE.read_bytes())['rows']
            for i, row in enumerate(rows, 1): row.update(list_sequence=i, capture_date='2025-01-08')
            with (folder/'journeys.csv').open('w', newline='', encoding='utf-8-sig') as f:
                w=csv.DictWriter(f, fieldnames=app.FIELDS); w.writeheader(); w.writerows(rows)
            (folder/'checks.json').write_text(json.dumps([{'month':'2025-01','expected_trip_count':3,'expected_distance_displayed':'123,4','status':'complete'}]), encoding='utf-8')
            app.build_reports(folder)
            self.assertTrue((folder/'journeys.xlsx').exists())
            data=json.loads((folder/'journeys.json').read_text(encoding='utf-8'))
            self.assertEqual(sum(int(x['grouped_trip_count']) for x in data['entries']),3)
            self.assertEqual(data['entries'][0]['resolved_date'],'2025-01-08')
            self.assertTrue((folder/'journeys.xml').exists())
            with zipfile.ZipFile(folder/'journeys.xlsx') as z:
                wb=ET.fromstring(z.read('xl/workbook.xml'))
                names=[s.attrib['name'] for s in wb.findall('.//{*}sheet')]
                self.assertEqual(names,['Monthly summary','2025-01'])
                sheet=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
                vals=[n.text for n in sheet.findall('.//{*}t')]
                self.assertIn('Displayed entry distance sum (km)', vals)
            self.assertEqual(app.load_validated_workbook(folder/'journeys.xlsx')['Monthly summary']['E2'].value,3)
            monthly=app.load_validated_workbook(folder/'journeys.xlsx')['2025-01']
            headers={cell.value:cell.column for cell in monthly[1]}
            self.assertEqual(monthly.auto_filter.ref.split(':')[0],'A1')
            self.assertEqual(monthly.freeze_panes,'A2')
            self.assertEqual(monthly.cell(4,headers['distance_km_numeric']).value,'=SUM(I2:I3)')
            self.assertEqual(monthly.cell(3,headers['grouped_trip_count']).value,2)
            self.assertEqual(monthly.cell(3,headers['grouped_trip_count']).fill.fgColor.rgb,'00FFF2CC')

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_report_only_cli_generates_offline_formats(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            row=app.parse_screen(FIXTURE.read_bytes())['rows'][0]
            row.update(list_sequence=1,capture_date='2025-01-08')
            app.atomic_csv(folder/'capture.csv',[row],app.FIELDS+app.PROVENANCE_FIELDS)
            app.atomic_json(folder/'checks.json',[{'month':'2025-01','status':'complete','expected_trip_count':1,'expected_distance_displayed':'12.3'}])
            result=subprocess.run([sys.executable,app.__file__,'--report-only',str(folder),'--formats','all'],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            for name in ('journeys.csv','journeys.json','journeys.xml','journeys.xlsx'):
                self.assertTrue((folder/name).exists())

    @unittest.skipUnless(HAS_OPENPYXL, 'XLSX tests require the declared openpyxl dependency')
    def test_default_cli_format_is_xlsx_only_and_formats_are_validated(self):
        self.assertEqual(app.parse_formats('xlsx,csv'), ['xlsx','csv'])
        self.assertEqual(app.parse_formats('all'), ['xlsx','csv','json','xml'])
        with self.assertRaises(argparse.ArgumentTypeError): app.parse_formats('xlsx,ods')
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp); row=app.parse_screen(FIXTURE.read_bytes())['rows'][0]
            app.atomic_csv(folder/'capture.csv',[row],app.FIELDS)
            app.atomic_json(folder/'checks.json',[{'month':'2025-01','status':'complete','expected_trip_count':1}])
            result=subprocess.run([sys.executable,app.__file__,'--report-only',str(folder)],capture_output=True,text=True)
            self.assertEqual(result.returncode,0,result.stderr)
            self.assertTrue((folder/'journeys.xlsx').exists())
            for name in ('journeys.csv','journeys.json','journeys.xml'):
                self.assertFalse((folder/name).exists(), name)
            self.assertEqual(json.loads((folder/'checks.json').read_text(encoding='utf-8'))[0]['status'],'complete')

    def test_capture_data_cleanup_preserves_unowned_files_and_retention_modes(self):
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'; data=out/'capture-data'; raw=data/'raw'; raw.mkdir(parents=True)
            (data/'.owner.json').write_text('{"kind":"bmw-journey-exporter-capture-data","owned_raw_files":["raw/screen-00001.xml"]}', encoding='utf-8')
            for name in ('capture-state.json','capture.csv','checks.json','manifest.json'):
                (data/name).write_text('owned', encoding='utf-8')
            (raw/'screen-00001.xml').write_text('<xml/>', encoding='utf-8')
            (data/'user-note.txt').write_text('preserve', encoding='utf-8')
            app.cleanup_capture_data(out)
            self.assertTrue((out/'journeys.xlsx').parent.exists())
            self.assertEqual((data/'user-note.txt').read_text(encoding='utf-8'),'preserve')
            self.assertFalse((data/'capture.csv').exists())
            self.assertFalse((raw/'screen-00001.xml').exists())
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'; data=out/'capture-data'; raw=data/'raw'; raw.mkdir(parents=True)
            (data/'.owner.json').write_text('{"kind":"bmw-journey-exporter-capture-data","owned_raw_files":["raw/screen-00001.xml"]}', encoding='utf-8')
            (data/'capture.csv').write_text('saved', encoding='utf-8'); (raw/'screen-00001.xml').write_text('<xml/>', encoding='utf-8')
            app.cleanup_capture_data(out,keep_capture=True)
            self.assertTrue((data/'capture.csv').exists()); self.assertFalse((raw/'screen-00001.xml').exists())
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'capture'; data=out/'capture-data'; raw=data/'raw'; raw.mkdir(parents=True)
            (data/'.owner.json').write_text('{"kind":"bmw-journey-exporter-capture-data","owned_raw_files":["raw/screen-00001.xml"]}', encoding='utf-8')
            (data/'capture.csv').write_text('saved', encoding='utf-8'); (raw/'screen-00001.xml').write_text('<xml/>', encoding='utf-8')
            app.cleanup_capture_data(out,keep_capture=True,keep_debug=True)
            self.assertTrue((raw/'screen-00001.xml').exists())

    def test_capture_main_cleanup_interruption_mismatch_and_report_failure(self):
        for outcome, expected_result, report_failure in (
            ('complete',0,False),('interrupt',1,False),('incomplete',2,False),('complete',1,True)):
            with self.subTest(outcome=outcome,report_failure=report_failure), tempfile.TemporaryDirectory() as tmp:
                out=Path(tmp)/'capture'
                seen={}
                class FakeExporter:
                    def __init__(self,args):
                        self.out=Path(args.output); self.out.mkdir()
                        self.data_dir=self.out/'capture-data'; (self.data_dir/'raw').mkdir(parents=True)
                        (self.data_dir/'.owner.json').write_text('{"kind":"bmw-journey-exporter-capture-data","owned_raw_files":["raw/screen-00001.xml"]}', encoding='utf-8')
                        (self.data_dir/'raw'/'screen-00001.xml').write_text('<xml/>', encoding='utf-8')
                        self.rows=[{'month':'2026-01','grouped_trip_count':1}] if outcome=='complete' else []; self.reports=[]; self.session={}; self.outcome=outcome
                    def checkpoint(self):
                        for name in ('capture-state.json','capture.csv','checks.json','manifest.json'):
                            (self.data_dir/name).write_text('checkpoint', encoding='utf-8')
                    def run(self):
                        if self.outcome=='interrupt': raise KeyboardInterrupt()
                        self.reports[0]['status']='incomplete' if self.outcome=='incomplete' else 'complete'
                        self.reports[0]['expected_trip_count']=1
                    def restore_keep_awake(self): pass
                def fake_reports(folder,replacements,formats):
                    seen['formats']=formats
                    if report_failure: raise OSError('simulated report failure')
                    (Path(folder)/'journeys.xlsx').write_text('validated output', encoding='utf-8')
                stdout=io.StringIO(); stderr=io.StringIO()
                with mock.patch.object(app.sys,'argv',['cli','--month','2026-01','--output',str(out)]), \
                     mock.patch.object(app,'adb_devices',return_value=[['MOCK','device']]), \
                     mock.patch.object(app,'Exporter',FakeExporter), \
                     mock.patch.object(app,'build_reports',side_effect=fake_reports), \
                     contextlib.redirect_stdout(stdout),contextlib.redirect_stderr(stderr):
                    result=app.main()
                self.assertEqual(result,expected_result)
                self.assertEqual(seen['formats'],['xlsx'])
                if outcome=='complete' and not report_failure:
                    self.assertTrue((out/'journeys.xlsx').exists())
                    self.assertFalse((out/'capture-data').exists())
                else:
                    self.assertTrue((out/'capture-data'/'capture-state.json').exists())
                    self.assertTrue((out/'capture-data'/'raw'/'screen-00001.xml').exists())
                    recovery_text=stdout.getvalue()+stderr.getvalue()
                    self.assertIn('make resume OUTPUT=',recovery_text)
                    self.assertIn('python3 bmw_export_journeys.py --resume',recovery_text)

    def test_keep_capture_and_keep_debug_cli_retention(self):
        for flag,raw_expected in (('--keep-capture',False),('--keep-debug',True)):
            with self.subTest(flag=flag), tempfile.TemporaryDirectory() as tmp:
                out=Path(tmp)/'capture'
                class FakeExporter:
                    def __init__(self,args):
                        self.out=Path(args.output); self.out.mkdir()
                        self.data_dir=self.out/'capture-data'; (self.data_dir/'raw').mkdir(parents=True)
                        (self.data_dir/'.owner.json').write_text('{"kind":"bmw-journey-exporter-capture-data","owned_raw_files":["raw/screen-00001.xml"]}', encoding='utf-8')
                        (self.data_dir/'raw'/'screen-00001.xml').write_text('<xml/>', encoding='utf-8')
                        self.rows=[{'month':'2026-01','grouped_trip_count':1}]; self.reports=[]; self.session={}
                    def checkpoint(self):
                        for name in ('capture-state.json','capture.csv','checks.json','manifest.json'):
                            (self.data_dir/name).write_text('checkpoint', encoding='utf-8')
                    def run(self):
                        self.reports[0]['status']='complete'; self.reports[0]['expected_trip_count']=1
                    def restore_keep_awake(self): pass
                with mock.patch.object(app.sys,'argv',['cli','--month','2026-01','--output',str(out),flag]), \
                     mock.patch.object(app,'adb_devices',return_value=[['MOCK','device']]), \
                     mock.patch.object(app,'Exporter',FakeExporter), \
                     mock.patch.object(app,'build_reports',side_effect=lambda folder,replacements,formats: (Path(folder)/'journeys.xlsx').write_text('ok', encoding='utf-8')) , \
                     contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(app.main(),0)
                data=out/'capture-data'
                self.assertTrue((data/'capture.csv').exists())
                self.assertEqual((data/'raw'/'screen-00001.xml').exists(),raw_expected)

if __name__ == '__main__': unittest.main()
